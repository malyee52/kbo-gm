"""진행 중이던 시즌(기본 2026)의 기록을 자료집 출처(D1: PsyproLEE/KBO_statics, KBO 공식 기록 수집본, 매일 갱신)의 최신본으로 바꾼다.

쓰는 법: python tools/refresh_season.py [시즌]   (그 뒤 python tools/build_data.py 로 public/data를 다시 만든다)

자료집(data-src/KBO_단장게임_자료집.xlsx)의 그 시즌 행만 고친다:
  - 타자시즌·투수시즌: 선수ID(KBO playerId)로 기록 열을 바꾼다. 도루·도루 실패는 출처의 선수별 연도 기록(careers.json)에서.
    출처에만 있는 선수(새로 1군에 오른 선수)는 행을 더한다.
  - 시즌순위: 경기·승·패·무·승률·홈·방문, 최종 순위(가을야구 전에는 정규시즌 순위)
  - 선수마스터: 새 선수 행을 더하고, 그 시즌이 마지막 1군 시즌이 된 선수를 고친다
수식 열(시즌 나이 등)은 그대로 둔다 (build_data.py는 기록 열만 읽는다). 2026-10-12 사용자: "26시즌은 사실상 끝난 상태"
"""
import io
import json
import re
import sys
import tempfile
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'data-src' / 'KBO_단장게임_자료집.xlsx'
BASE = 'https://raw.githubusercontent.com/PsyproLEE/KBO_statics/main/web/public/data/'
SEASON = int(sys.argv[1]) if len(sys.argv) > 1 else 2026


def fetch(name, tmp):
    p = Path(tmp) / name
    urllib.request.urlretrieve(BASE + name, p)
    return json.load(io.open(p, encoding='utf-8'))


def num(v):
    if v in (None, '', '-'):
        return 0
    try:
        return int(v)
    except ValueError:
        return float(v)


def outs_of(ip):
    """'172 1/3' → 517"""
    m = re.match(r'^\s*(\d+)?\s*(?:(\d)/3)?\s*$', str(ip or '0'))
    if not m:
        return 0
    return int(m.group(1) or 0) * 3 + int(m.group(2) or 0)


def header(ws):
    return [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]


def main():
    with tempfile.TemporaryDirectory() as tmp:
        meta = fetch('meta.json', tmp)
        if int(meta.get('season', 0)) != SEASON:
            sys.exit(f'출처의 현재 시즌이 {meta.get("season")}입니다 ({SEASON}이 아님)')
        hitters = fetch('hitters.json', tmp)
        pitchers = fetch('pitchers.json', tmp)
        standings = fetch('standings.json', tmp)
        careers = fetch('careers.json', tmp)
        profiles = fetch('players.json', tmp)
    print(f'출처 갱신 시각: {meta.get("updatedAt")}')

    wb = load_workbook(SRC)
    ws_bat, ws_pit, ws_st, ws_m = wb['타자시즌'], wb['투수시즌'], wb['시즌순위'], wb['선수마스터']
    hb, hp, hs, hm = header(ws_bat), header(ws_pit), header(ws_st), header(ws_m)
    col = lambda h, name: h.index(name) + 1

    # 구단 → 프랜차이즈 (그 시즌 순위표)
    franchise = {}
    for r in ws_st.iter_rows(min_row=2, values_only=True):
        if r[0] == SEASON:
            franchise[r[hs.index('팀')]] = r[hs.index('프랜차이즈')]

    master = {}
    for i, r in enumerate(ws_m.iter_rows(min_row=2, values_only=True), start=2):
        master[str(r[0])] = i

    def career_row(pid, kind):
        c = careers.get(pid, {}).get(kind)
        if not c:
            return None
        rows = [x for x in c['rows'] if str(x.get('연도')) == str(SEASON)]
        return rows[-1] if rows else None

    def update_rows(ws, h, src, fill):
        rows = {}
        for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if r[0] == SEASON:
                rows.setdefault(str(r[h.index('선수ID')]), []).append(i)
        seen = set()
        changed = added = 0
        for s in src:
            pid = str(s['playerId'])
            seen.add(pid)
            targets = rows.get(pid)
            if targets:
                i = targets[0]
            else:
                i = ws.max_row + 1
                added += 1
            vals = fill(s, pid, i, bool(targets))
            for name, v in vals.items():
                ws.cell(row=i, column=col(h, name), value=v)
            changed += 1
        stale = [pid for pid in rows if pid not in seen]
        return changed, added, stale

    def new_player(s, pid, kind):
        """선수마스터에 없는 선수: 출처 프로필로 최소한의 행을 더한다"""
        if pid in master:
            i = master[pid]
            ws_m.cell(row=i, column=col(hm, '마지막 1군 시즌'), value=SEASON)
            ws_m.cell(row=i, column=col(hm, '2026 기록 있음'), value='Y')
            return
        pr = profiles.get(pid, {})
        birth = pr.get('birth', '')
        m = re.match(r'(\d{4})년 (\d{2})월 (\d{2})일', birth)
        pos = pr.get('position', '')
        hand = re.search(r'\((.+)\)', pos)
        i = ws_m.max_row + 1
        vals = {
            '선수ID': pid, '선수명': s['선수명'], '구분': '투수' if kind == 'pitcher' else '타자',
            '주포지션': re.sub(r'\(.*\)', '', pos) or None, '외국인': 'N', '외국인 판별 근거': '출처 프로필 (refresh_season.py)',
            '생년월일': f'{m.group(1)}-{m.group(2)}-{m.group(3)}' if m else None, '출생연도': int(m.group(1)) if m else None,
            '투타': hand.group(1) if hand else None, '입단연도(추정)': SEASON, '입단연도 근거': '첫 1군 시즌',
            '첫 1군 시즌': SEASON, '마지막 1군 시즌': SEASON, '1군 시즌 수': 1, '첫 소속팀': s['팀명'], '마지막 소속팀': s['팀명'],
            '소속팀 이력': s['팀명'], '2026 기록 있음': 'Y', '출신교·경력(2026 프로필)': pr.get('career'),
            '신장/체중(2026 프로필)': pr.get('body'), '보조자료 연결': 'N',
        }
        for name, v in vals.items():
            ws_m.cell(row=i, column=col(hm, name), value=v)
        master[pid] = i
        print(f'  새 선수: {s["선수명"]} ({pid}, {s["팀명"]})')

    def bat_fill(s, pid, i, exists):
        cr = career_row(pid, 'hitter') or {}
        v = {
            'G': num(s['G']), 'PA': num(s['PA']), 'AB': num(s['AB']), 'R': num(s['R']), 'H': num(s['H']), '2B': num(s['2B']),
            '3B': num(s['3B']), 'HR': num(s['HR']), 'TB': num(s['TB']), 'RBI': num(s['RBI']), 'SB': num(cr.get('SB')), 'CS': num(cr.get('CS')),
            'BB': num(s['BB']), 'IBB': num(s['IBB']), 'HBP': num(s['HBP']), 'SO': num(s['SO']), 'GDP': num(s['GDP']),
            'SAC': num(s['SAC']), 'SF': num(s['SF']), 'AVG': num(s['AVG']), 'OBP': num(s['OBP']), 'SLG': num(s['SLG']), 'OPS': num(s['OPS']),
        }
        if not exists:
            new_player(s, pid, 'hitter')
            pr = profiles.get(pid, {})
            v.update({'시즌': SEASON, '팀': s['팀명'], '프랜차이즈': franchise.get(s['팀명'], s['팀명']), '선수ID': pid, '선수명': s['선수명'],
                      '구분': '타자', '포지션(보조)': re.sub(r'\(.*\)', '', pr.get('position', '')) or None, '외국인': 'N',
                      '출생연도': ws_m.cell(row=master[pid], column=col(hm, '출생연도')).value,
                      '도루 출처': 'KBO 공식', '볼넷·삼진 등 출처': 'KBO 공식', '보조자료 연결': '없음'})
        else:
            v['팀'] = s['팀명']
            v['프랜차이즈'] = franchise.get(s['팀명'], s['팀명'])
        return v

    def pit_fill(s, pid, i, exists):
        outs = outs_of(s['IP'])
        v = {
            'G': num(s['G']), 'W': num(s['W']), 'L': num(s['L']), 'SV': num(s['SV']), 'HLD': num(s['HLD']), 'IP(표기)': str(s['IP']),
            '아웃카운트': outs, 'TBF': num(s['TBF']), 'H': num(s['H']), 'HR': num(s['HR']), 'BB': num(s['BB']), 'IBB': num(s['IBB']),
            'HBP': num(s['HBP']), 'SO': num(s['SO']), 'R': num(s['R']), 'ER': num(s['ER']),
            'ERA': round(num(s['ER']) * 27 / outs, 2) if outs else None,
            'WHIP(계산)': round((num(s['H']) + num(s['BB'])) * 3 / outs, 2) if outs else None,
            'CG': num(s['CG']), 'SHO': num(s['SHO']), 'QS': num(s['QS']), 'BSV': num(s['BSV']), 'WP': num(s['WP']), 'BK': num(s['BK']),
            '투구수': num(s['NP']),
        }
        if not exists:
            new_player(s, pid, 'pitcher')
            v.update({'시즌': SEASON, '팀': s['팀명'], '프랜차이즈': franchise.get(s['팀명'], s['팀명']), '선수ID': pid, '선수명': s['선수명'],
                      '외국인': 'N', '출생연도': ws_m.cell(row=master[pid], column=col(hm, '출생연도')).value, '보조자료 연결': '없음'})
        else:
            v['팀'] = s['팀명']
            v['프랜차이즈'] = franchise.get(s['팀명'], s['팀명'])
        return v

    n, a, stale = update_rows(ws_bat, hb, hitters, bat_fill)
    print(f'타자: {n}명 갱신 (새 행 {a}){", 출처에 없는 기존 행 " + str(stale) if stale else ""}')
    n, a, stale = update_rows(ws_pit, hp, pitchers, pit_fill)
    print(f'투수: {n}명 갱신 (새 행 {a}){", 출처에 없는 기존 행 " + str(stale) if stale else ""}')

    by_team = {s['팀명']: s for s in standings}
    for i, r in enumerate(ws_st.iter_rows(min_row=2, values_only=True), start=2):
        if r[0] != SEASON:
            continue
        s = by_team.get(r[hs.index('팀')])
        if not s:
            continue
        for name, key in (('경기', '경기'), ('승', '승'), ('패', '패'), ('무', '무')):
            ws_st.cell(row=i, column=col(hs, name), value=num(s[key]))
        ws_st.cell(row=i, column=col(hs, '승률(KBO 표기)'), value=num(s['승률']))
        ws_st.cell(row=i, column=col(hs, '최종 순위(KBO 표기)'), value=num(s['순위']))
        ws_st.cell(row=i, column=col(hs, '홈(승-무-패)'), value=s['홈'])
        ws_st.cell(row=i, column=col(hs, '방문(승-무-패)'), value=s['방문'])
    print('순위:', ', '.join(f'{s["순위"]}.{s["팀명"]} {s["경기"]}경기 {s["승"]}-{s["패"]}-{s["무"]}' for s in standings))

    wb.save(SRC)
    print(f'저장: {SRC}')


if __name__ == '__main__':
    main()
