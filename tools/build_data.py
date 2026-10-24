#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""자료집 엑셀 → 게임 데이터 JSON.

사용법:  python3 tools/build_data.py [엑셀 경로] [출력 폴더]
기본값:  data-src/KBO_단장게임_자료집.xlsx → public/data/

만드는 파일
  players.json          선수 마스터 (선수 한 명당 한 줄)
  seasons/<연도>.json   그 해의 구단, 리그 합계, 타자·투수 시즌 기록
  meta.json             연도 목록, 구단 계보, 연도별 규칙, 능력 산출용 사전 평균
  contracts.json        FA·비FA 다년 계약 (2000~2025 계약 시작, 금액 미검증 보조자료)
  drafts.json           입단 연도별 신인 지명 (1982~2027, 원년 멤버·육성선수 포함)

함께 읽는 파일
  (연도별 규칙)                 meta.json의 rules는 rules_for·innings_rules에 적은 연표 (자료집 '제도연표'와 조사 결과를 옮긴 값).
                                 이것만 다시 만들려면: build_data.py --rules-only (rule_changes.json도 함께, 다른 JSON과 저장 호환은 그대로)
  data-src/외국인_개막명단.csv   시작 연도 개막 때 구단별 외국인과 아시아쿼터 여부 (meta.json의 foreignOpening, 선수 마스터의 asia).
                                연봉(달러)·계약금(달러)·금액 출처 열이 있으면 시작 계약에 그대로 쓴다 (없으면 게임 생성식)
                                 자료집에 날짜·국적이 없어 따로 둔다. '확인' 열이 '미확인'이면 추정값이다 ('확인'·'출처 확인'은 확인값).
                                 구분: 개막 명단(일반·아시아쿼터) 또는 제외(시즌 중 교체 선수). 계약구분: 아시아쿼터면 시장에서도 아시아쿼터.
                                 기록소속: 개막 구단과 그 해 기록 소속이 다를 때(시즌 중 이적) 기록 소속. 월드는 개막 구단에서 시작한다.
  data-src/제도연표_보정.csv      자료집 '제도연표'의 미확인 항목을 위키로 확인한 값과 출처 (2026-10-23 사용자 결정 "위키 확인도 출처 확인").
                                 엑셀은 고치지 않고 이 파일에 남긴다. 읽는 코드는 없고, 반영한 곳은 '게임 반영' 열 (rules_for 등).
  data-src/병역_면제명단.csv     국제대회 금메달 대표팀과 2026 아시안게임 새 면제자 (military.json). 출처 열 참조.
                                 상무·경찰 복무는 자료집 선수마스터의 '출신교·경력(2026 프로필)'에서 읽는다.
                                 엑셀의 다른 내용을 바꾸지 않고 이것만 다시 만들려면: build_data.py --military-only
  data-src/baseballchart/kbo_draft_1982-2027.csv
                                 입단 연도별 실제 신인 지명 (drafts.json). baseballchart.kr에서 받은 원본, CC BY-NC-SA 2.0 KR.
                                 같은 폴더의 SOURCE.md 참조. 지명받고 그해 입단하지 않은 선수(해외 진출 등)는 실제 입단 해로 옮긴다 (defer_unsigned).
                                 이것만 다시 만들려면: build_data.py --drafts-only
  data-src/연고지.csv            구단 계보별 연고 도시·위도·경도 (homes.json, 원정 거리 컨디션). 출처 열 참조, 모두 임시값.
                                 이것만 다시 만들려면: build_data.py --homes-only
  data-src/감독.csv              연도별 구단 감독 실명 (managers.json). 한국어 위키백과 구단별 감독 틀·감독 목록, 2026 보도로 만든 것.
                                 '시점'이 개막이면 그 시즌을 시작한 감독, 중도면 시즌 중 바뀐 감독. 메모에 확인할 점.
                                 감독의 운영 경향(도루 시도·희생번트·구원 등판)은 그 감독이 시즌 내내 맡은 해의 팀 기록으로 계산한다.
                                 이것만 다시 만들려면: build_data.py --managers-only
  data-src/규칙변경.csv          연도별 리그 규칙 변경 (rule_changes.json, 개막 "올해 달라지는 규칙" 알림). 열: 연도, 날짜(MM-DD), 분류, 키,
                                 제목, 전, 후, 단위, 반영, 확인, 출처, 비고("게임은 단순화: …"로 시작하면 단순화 문구). 키가 있으면 전·후는 게임 값에서 읽는다.
  data-src/공인구.csv            공인구 반발계수 기준 범위·단일/복수·검사 평균 (rule_changes.json의 ball). 확인이 조사값이면 웹 조사 값.
                                 둘만 다시 만들려면: build_data.py --rules-only
  data-src/hardhit_derived/      하드힛(https://hardhit.ai, 비상업적 이용) 기록을 게임 선수 ID에 연결한 가공값.
                                 원본(data-src/hardhit/, 커밋 안 함)에서 tools/hardhit_derive.py가 만든다. 원본이 없어도 이 폴더로 빌드된다.
                                 bat_running.csv: 자료집에 비어 있는 타자 도루·도루실패·병살타를 채운다
                                 자료집 값과 다르면 하드힛 값을 쓴다 (2026-10-22 사용자 결정, 모든 하드힛 가공값에 같은 규칙)
                                 pitch_hand.csv: 투타가 빈 투수의 던지는 손·언더핸드를 채운다 (여러 해 값이 갈리면 가장 많은 쪽)
                                 park.csv: 구단·연도별 홈 구장 계수 (meta.json의 parkFactors, park_factors 참조)
                                 park_components.csv: 구장별 구성 지수 → 사건별 구장 배율 (meta.json의 parkEvents, park_events 참조)
  data-src/wiki/                 위키데이터·위키백과·나무위키에서 뽑은 가공값 (tools/wiki/ 하네스, data-src/wiki/README.md).
                                 빈 칸만 채운다 (2026-10-23 사용자 결정, 원자료 > 위키백과 > 나무위키). 우리 값과 다른 칸은 바꾸지 않는다.
                                 인적사항.csv: 생년·투타·언더핸드·세부 포지션(빈 칸·IF·OF만), 병역(실제 자료가 없는 선수, military.json src 'wiki')
                                 기록_채움후보.csv: 있는 시즌 행의 빈 칸만 (빈 행 '(행 없음)'은 더하지 않는다: 선수 구성이 바뀐다)
                                 키·몸무게는 아직 게임에 쓰지 않아 싣지 않는다. 결과는 reports/wiki-fill.md
                                 세부 포지션은 1군 마지막 해가 2000년까지인 야수만 (그 뒤 선수는 '세부 포지션(위키)' 칸으로만 남는다)
  data-src/수상.csv · 올스타.csv  골든글러브·MVP·신인왕·한국시리즈 MVP·올스타전 MVP와 올스타전 경기 (awards.json, allstar_history.json).
                                 나무위키 덤프에서 tools/wiki/s6_awards.py가 만든다. 짝 결과는 reports/wiki-awards.md
  data-src/구장.csv · 구장_사용.csv  구장 × 제원 기간(펜스·좌석·잔디·돔)과 구단 계보별 주 구장 (parks.json). 나무위키 구장 문서에서 읽어 적은 값.
                                 열은 docs/구장효과_기획안.md 7.1·7.2 (plan/park-effects). 이것만 다시 만들려면: build_data.py --history-only
  data-src/수비이닝.csv · 수비평판.csv  KBO 공식 사이트 수비 기록 2001~2026 (사용자 수집, 선수·연도·구단·자리별 G·GS·이닝·E·PO·A·DP·PB·SB·CS)과
                                 나무위키·위키백과 서술의 수비 평판 (등급·시기·문서 제목만) → defense.json. 저장소 밖 원본에서 tools/build_defense.py가 만든다.
                                 이것만 다시 만들려면: build_data.py --defense-only (능력치 v2 16-3a, 판정 함수는 src/engine/fielding.ts)
  data-src/구종.csv · 위키_구종.csv  visualbaseball.com 구종 자료 2016~ (사용자 수집, 선수·연도·구종별 구속·무브·릴리스·낙차·사용률·투구수·RV/100)와
                                 나무위키·위키백과 덤프의 구종 목록 (낱말·순서·주무기·시기·최고 구속만) → pitches.json (+ 15.5 측정 표 prior).
                                 저장소 밖 원본에서 tools/build_pitches.py가 만든다. 이것만 다시 만들려면: build_data.py --pitches-only (능력치 v2 16-1a·16-1b)
"""
import csv
import json
import math
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
ARGS = [a for a in sys.argv[1:] if not a.startswith('--')]
SRC = Path(ARGS[0]) if ARGS else ROOT / 'data-src' / 'KBO_단장게임_자료집.xlsx'
HOMES_CSV = ROOT / 'data-src' / '연고지.csv'
PBP_DERIVED = ROOT / 'data-src' / 'pbp_derived'
HARDHIT_DERIVED = ROOT / 'data-src' / 'hardhit_derived'
FOREIGN_OPENING = ROOT / 'data-src' / '외국인_개막명단.csv'
TRAITS_CSV = ROOT / 'data-src' / '특수능력.csv'
MILITARY_CSV = ROOT / 'data-src' / '병역_면제명단.csv'
MANAGERS_CSV = ROOT / 'data-src' / '감독.csv'
RULES_CSV = ROOT / 'data-src' / '규칙변경.csv'
BALL_CSV = ROOT / 'data-src' / '공인구.csv'
HARDHIT = ROOT / 'data-src' / 'hardhit_derived'
HARDHIT_REPORT = ROOT / 'reports' / 'hardhit-fill.md'
# 그 해 1군 기록이 자료에 없지만 선수단에 넣어야 하는 선수 (예: 2026 김광현). 기록 0인 행을 그 시즌에 더한다
EXTRA_ROSTER = ROOT / 'data-src' / '추가_명단.csv'
WIKI = ROOT / 'data-src' / 'wiki'
WIKI_REPORT = ROOT / 'reports' / 'wiki-fill.md'
# --no-wiki: 위키 가공값을 넣지 않은 기반 자료를 만든다 (tools/wiki/s5_merge.py가 '우리 값'으로 견줄 때 쓴다.
# 위키 값이 이미 들어간 public/data와 견주면 채운 칸이 채움 후보에서 빠져 다음 빌드에서 사라진다)
NO_WIKI = '--no-wiki' in sys.argv
# 숨겨진 특수능력: CSV의 한글 이름 → 엔진 코드 (src/engine/traits.ts의 TRAIT_INFO와 같아야 한다)
TRAIT_LABELS = {
    '승부사': 'clutch', '가을 사나이': 'october', '위기 탈출': 'escape', '철인': 'ironman', '늦게 지는 꽃': 'evergreen',
    '에이스 체력': 'stamina', '안방 사나이': 'homer', '좌완 킬러': 'vsLeft', '우완 킬러': 'vsRight',
}


def trait_key(pid):
    """선수 ID → traits.json의 키. 소금을 붙인 FNV-1a 32비트 해시 (src/engine/traits.ts의 traitKey와 같은 값)"""
    h = 0x811c9dc5
    for b in f'kbo-gm/traits/{pid}'.encode('utf-8'):
        h ^= b
        h = (h * 0x01000193) & 0xffffffff
    return f'{h:08x}'


def build_traits(known_ids):
    """data-src/특수능력.csv → {해시: [코드, ...]}. 모르는 선수 ID나 능력 이름은 오류"""
    table = {}
    if not TRAITS_CSV.exists():
        return table
    with TRAITS_CSV.open(encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(f):
            pid = str(r['선수ID']).strip()
            codes = [TRAIT_LABELS.get(x.strip()) for x in (r['능력'] or '').split(';') if x.strip()]
            if not codes:
                continue
            if pid not in known_ids:
                raise SystemExit(f'특수능력.csv: 없는 선수 ID {pid} ({r.get("선수명")})')
            if None in codes:
                raise SystemExit(f'특수능력.csv: 모르는 능력 이름 "{r["능력"]}" ({r.get("선수명")}). 가능한 값: {", ".join(TRAIT_LABELS)}')
            table[trait_key(pid)] = sorted(set(codes))
    return dict(sorted(table.items()))


def dump_traits(table):
    (OUT).mkdir(parents=True, exist_ok=True)
    (OUT / 'traits.json').write_text(json.dumps({
        'version': 1,
        'note': '숨겨진 특수능력. 키는 선수 ID의 해시(tools/build_data.py trait_key), 값은 능력 코드. 화면에 보이지 않는다.',
        'traits': table,
    }, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'특수능력: {len(table)}명 → {OUT / "traits.json"}')
DRAFTS = ROOT / 'data-src' / 'baseballchart' / 'kbo_draft_1982-2027.csv'
DRAFT_POS = {'투수': 'P', '포수': 'C', '내야수': 'IF', '외야수': 'OF', '지명타자': 'DH'}


def draft_pos(text):
    """포지션 표기 → (P·C·IF·OF·DH, 투타겸업 여부). '투수/외야수'처럼 둘이면 앞의 것. 투타겸업은 투수로 본다"""
    if not text:
        return 'IF', False
    if text == '투타겸업':
        return 'P', True
    first = re.split(r'[/,·]', text)[0].strip()
    return DRAFT_POS.get(first, 'IF'), len(re.split(r'[/,·]', text)) > 1


def draft_rows():
    """baseballchart.kr 원본 → 입단 연도별 지명 목록 (원본 순서 그대로)"""
    out = defaultdict(list)
    if not DRAFTS.exists():
        return out
    with DRAFTS.open(encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(f):
            kind = r['갈래'].strip()
            rnd = None
            m = re.fullmatch(r'(\d+)라운드', kind)
            if m:
                kind, rnd = '라운드', int(m.group(1))
            overall = re.search(r'(\d+)순위', r['자리'] or '')
            pos, two_way = draft_pos(r['포지션'].strip())
            career = [x.strip() for x in (r['이력'] or '').split('-') if x.strip()]
            name = re.sub(r'\s*\(.*\)$', '', r['선수']).strip()  # '김바위 (김용윤)' → 김바위
            row = {
                'kind': kind, 'round': rnd, 'overall': int(overall.group(1)) if overall else None,
                'team': r['지명 구단'].strip(), 'name': name, 'pos': pos,
                'school': career[0] if career else '',
                # 이력에 대학(…대, …대학교)이 있으면 대졸로 본다
                'univ': any(x.endswith('대') or '대학' in x for x in career),
                'games': int(r['통산 경기'] or 0),
            }
            if two_way:
                row['twoWay'] = True
            out[str(r['연도'])].append(row)
    return out
OUT = Path(ARGS[1]) if len(ARGS) > 1 else ROOT / 'public' / 'data'


def master_games(out_dir):
    """선수 ID → 실제 KBO 1군 통산 경기 (시즌마다 타자·투수 기록 중 많은 쪽)"""
    meta = json.loads((out_dir / 'meta.json').read_text(encoding='utf-8'))
    total = defaultdict(int)
    for y in meta['years']:
        sd = json.loads((out_dir / 'seasons' / f'{y}.json').read_text(encoding='utf-8'))
        bi = sd['batFields'].index('g')
        pi = sd['pitFields'].index('g')
        g = defaultdict(int)
        for r in sd['bat']:
            g[r[0]] = max(g[r[0]], r[bi] or 0)
        for r in sd['pit']:
            g[r[0]] = max(g[r[0]], r[pi] or 0)
        for k, v in g.items():
            total[k] += v
    return total


def school_key(name):
    """'광주제일고'·'광주일고'처럼 같은 학교의 다른 표기를 맞춘다"""
    return (name or '').replace('제일', '일').replace('상업고', '상고').replace('공업고', '공고').strip()


def defer_unsigned(drafts, players, games_of):
    """지명 행을 실제 선수와 잇고, 지명받고 그해 입단하지 않은 행을 정리한다
    (2026-10-12 사용자: 2001년 채태인·추신수 같은 선수가 그해 드래프트에 나온다).
    해외 진출·대학 진학으로 지명만 받고 몇 해 뒤 입단한 선수가 지명 해 게임 드래프트에 나오면 실제보다 일찍 KBO에서 뛰고,
    게임은 지명 행을 선수 마스터와 '이름 + 입단 연도 ±1'로만 맞춰 가짜 동명 선수가 생겼다.
    - 잇기: 같은 이름의 국내 선수마다, 데뷔 해까지의 지명 가운데 통산 경기(원본 '통산 경기')가 그 선수의 실제 1군 통산 경기와
      맞는 행(±25%·5경기)을 그 선수의 입단 지명으로 보고 pid를 붙인다 (나이가 맞는 행만). 여럿이면 경기 수가 가장 가까운 행, 같으면 입단 해와 가까운 행.
      실제 입단(entryYear)이 지명 2년 넘게 뒤면 행을 입단 해로 옮기고 drafted에 원래 지명 연도를 남긴다 (예: 추신수 2007 SK 해외 → 2021).
    - 빼기: 이은 행보다 앞선 같은 이름·같은 출신교의 0경기 지명은 그 사람이 입단하지 않은 지명이다 (예: 채태인 2001 두산)
    돌려주는 값: (정리한 지명, 옮긴 목록, 뺀 목록)"""
    by_name = defaultdict(list)
    for m in players:
        if not m.get('foreign') and m.get('first'):
            by_name[m['name']].append(m)
    rows_of = defaultdict(list)
    for y, lst in drafts.items():
        for d in lst:
            if d['kind'] not in ('육성선수', '원년 멤버'):
                rows_of[d['name']].append((int(y), d))
    claim = {}
    drop = set()
    for name, rows_ in rows_of.items():
        rows_.sort(key=lambda x: x[0])
        taken = set()
        for m in sorted(by_name.get(name, []), key=lambda m: (m['first'], m['id'])):
            b = m.get('birthYear')
            gm = games_of.get(m['id'], 0)
            age_ok = lambda y: (16 <= y - b <= 45) if b else (y >= m['first'] - 20)
            near = lambda g: g > 0 and abs(g - gm) <= max(5, gm * 0.25)
            ok = [(y, d) for y, d in rows_ if id(d) not in taken and y <= m['first'] and age_ok(y) and near(d['games'])]
            if not ok:
                continue
            # 경기 수가 가장 가까운 행, 같으면 실제 입단 해와 가까운 행 (동명이인끼리 경기 수가 적으면 헷갈린다)
            y, d = min(ok, key=lambda x: (abs(x[1]['games'] - gm), abs(x[0] - (m.get('entryYear') or x[0])), -x[0]))
            taken.add(id(d))
            claim[id(d)] = m
            for y2, d2 in rows_:
                if y2 < y and id(d2) not in taken and d2['games'] == 0 and age_ok(y2) and school_key(d2['school']) == school_key(d['school']):
                    drop.add(id(d2))
                    taken.add(id(d2))
    out = defaultdict(list)
    moved, dropped = [], []
    for y in sorted(drafts, key=int):
        Y = int(y)
        for d in drafts[y]:
            if id(d) in drop:
                dropped.append(f"{d['name']} {Y}")
                continue
            m = claim.get(id(d))
            if not m:
                out[y].append(d)
                continue
            row = dict(d, pid=m['id'])
            e = m.get('entryYear') or Y
            if e > Y + 1:
                row['drafted'] = Y
                out[str(e)].append(row)
                moved.append(f"{d['name']} {Y}→{e}")
            else:
                out[y].append(row)
    return dict(out), moved, dropped


def dump_drafts(drafts, players):
    drafts, moved, dropped = defer_unsigned(drafts, players, master_games(OUT))
    (OUT / 'drafts.json').write_text(json.dumps({
        'version': 1, 'source': 'baseballchart.kr/draft (나무위키 「(연도) KBO 신인 드래프트」 문서 정리, 2026-09-21 수집)',
        'license': 'CC BY-NC-SA 2.0 KR', 'years': {y: drafts[y] for y in sorted(drafts, key=int)},
    }, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'신인 지명: 그해 입단하지 않은 지명 {len(moved)}건을 실제 입단 해로 옮김, {len(dropped)}건은 나중 지명이 있어 뺌')
    return drafts

POS = {'포수': 'C', '1루수': '1B', '2루수': '2B', '3루수': '3B', '유격수': 'SS', '좌익수': 'LF', '중견수': 'CF',
       '우익수': 'RF', '지명타자': 'DH', '내야수': 'IF', '외야수': 'OF', '선발투수': 'SP', '구원투수': 'RP',
       '마무리투수': 'CL', '투수': 'P'}
HAND = {'우': 'R', '좌': 'L', '양': 'S'}

# posSrc: 포지션 출처. 'H' 하드힛 확인값, 'E' 하드힛 해의 추정값(2009·파일 없는 타자), null 자료집 그대로 (fill_positions)
BAT_FIELDS = ['id', 'team', 'pos', 'g', 'pa', 'ab', 'r', 'h', 'd', 't', 'hr', 'rbi', 'sb', 'cs', 'bb', 'hbp', 'so', 'gdp', 'sf', 'sac', 'posSrc']
BAT_COLS = ['선수ID', '팀', '포지션(보조)', 'G', 'PA', 'AB', 'R', 'H', '2B', '3B', 'HR', 'RBI', 'SB', 'CS', 'BB', 'HBP', 'SO', 'GDP', 'SF', 'SAC']
# gmli: 하드힛 투수 situation의 등판 때 평균 레버리지 (2010~, 없으면 null). 시작 불펜 서열에 쓴다 (2026-10-22)
PIT_FIELDS = ['id', 'team', 'role', 'g', 'gs', 'w', 'l', 'sv', 'hld', 'outs', 'tbf', 'h', 'hr', 'bb', 'hbp', 'so', 'r', 'er', 'gmli']
PIT_COLS = ['선수ID', '팀', '역할(보조)', 'G', 'GS(보조)', 'W', 'L', 'SV', 'HLD', '아웃카운트', 'TBF', 'H', 'HR', 'BB', 'HBP', 'SO', 'R', 'ER']


def rows(ws):
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    for r in it:
        if r[0] is None:
            continue
        yield dict(zip(header, r))


def hands(bt):
    """'우투좌타', '우언투우타' → (던지는 손, 치는 손, 언더핸드 여부)"""
    if not bt or '투' not in bt:
        return None, None, False
    th, ba = bt.split('투')
    return HAND.get(th[0]), HAND.get(ba[0]), '언' in th


def school_of(career):
    """2026 프로필의 출신교·경력 문자열 → 'UNIV'(대졸) / 'HS'(고졸) / None.
    괄호 안 학교는 프로 입단 뒤 다닌 곳(사이버대 등)이라 빼고, 대학('...대')을 거쳤으면 대졸로 본다."""
    if not career:
        return None
    parts = [x.strip() for x in career.split('-') if x.strip() and not x.strip().startswith('(')]
    if any(x.endswith('대') or '대학' in x for x in parts):
        return 'UNIV'
    if any(x.endswith('고') for x in parts):
        return 'HS'
    return None


def contract_years(v):
    """계약 기간. '2(+2 선수옵션)'처럼 옵션이 붙으면 보장 기간(앞 숫자)만 쓴다."""
    if isinstance(v, int):
        return v
    m = re.match(r'\s*(\d+)', str(v or ''))
    return int(m.group(1)) if m else None


def rules_for(year):
    """연도별 규칙. source: 'confirmed' = 자료집 제도연표 또는 data-src/제도연표_보정.csv에서 출처 확인, 'assumed' = 미확인이라 가정한 값.
    2013년까지 1군 엔트리 26명은 그대로 가정값 (위키 조사의 '1982년 25명 → 1990년대 중반 26명'은 원문 문장을 찾지 못함, 보정 CSV 13)."""
    if year >= 2026: roster, rs = 29, 'confirmed'
    elif year >= 2020: roster, rs = 28, 'confirmed'
    elif year >= 2015: roster, rs = 27, 'confirmed'
    elif year == 2014: roster, rs = 26, 'confirmed'
    else: roster, rs = 26, 'assumed'
    if year < 1998: foreign, fs = 0, 'confirmed'
    elif year <= 2000: foreign, fs = 2, 'confirmed'
    elif year <= 2002: foreign, fs = 3, 'confirmed'
    elif year <= 2013: foreign, fs = 2, 'confirmed'  # 2003년 2명 보유로 환원 (보정 CSV 16, 나무위키 외국인 선수/KBO 리그)
    elif year <= 2025: foreign, fs = 3, 'confirmed'
    else: foreign, fs = 4, 'confirmed'
    # 한 경기 출전 외국인 (보유와 따로, 리그 규칙 기획안 D16): 2001~2002·2014~2019는 3명 보유·2명 출전 (자료집 제도연표 확인).
    # 그 밖의 해는 보유한 외국인이 모두 출전할 수 있다
    if 2001 <= year <= 2002 or 2014 <= year <= 2019: active, acs = 2, 'confirmed'
    else: active, acs = foreign, fs
    if year <= 1988: post = 2
    elif year <= 2014: post = 4
    else: post = 5
    rules = {
        'rosterSize': roster, 'rosterSizeSource': rs,
        'foreignLimit': foreign, 'foreignLimitSource': fs,
        'foreignActive': active, 'foreignActiveSource': acs,
        **innings_rules(year),
        'postseasonTeams': post, 'postseasonTeamsSource': 'confirmed' if year >= 2015 else 'assumed',
    }
    return rules


def innings_rules(year):
    """연장·무승부·승률 연표 (경기 엔진 2 14단계, docs/경기엔진2_기획안.md 3.10.1, 자료집 '제도연표' 연장 줄).
    근거: 나무위키 '연장전'(KBO 리그 절)·'승률', 잼쨩 웹 조사(2026-10-22, 기획안 3.10.1)와 빈 자료 조사(2026-10-23, 2_제도연표_미확인_조사결과.csv 102).
    위키로 확인한 제도는 '출처 확인'으로 본다 (사용자 결정 2026-10-23). 출처가 갈리는 해는 다수 설을 넣고 source를 'disputed'로 남긴다.
    시간제한(1982~2002 야간 4시간, 2004 4시간)은 엔진에 경기 시각이 없어 여기 넣지 않고, 엔진 params.timeLimit(연장 이닝마다 끝날 확률)으로 흉내 낸다.
      maxInnings: 1982~2002 15회, 2003~2007 12회 (2003은 출처 갈림: 빈 자료 조사는 2004년부터 12회), 2008 무제한(0, 무승부 없음),
        2009~2024 12회, 2025~ 11회
      nineInningsDhMonday: 2020 더블헤더·월요일 편성 경기는 9회 / nineInningsFrom: 2021 후반기(8월 10일~) 모든 경기 9회 (월×100 + 일)
      postseasonMaxInnings: 15회. 2003·2004 12회 (2003 출처 갈림: 15회 설), 2008 무제한 (출처 갈림: 15회 설)
      winPct: 'wl' 승 ÷ (승+패), 1987~1997 'half' 무승부 0.5승, 2003~2004 'wins' 다승제, 2009~2010 'tieLoss' 무승부 = 패"""
    if year <= 2002: inn = 15
    elif year == 2008: inn = 0
    elif year <= 2024: inn = 12
    else: inn = 11
    post = 0 if year == 2008 else 12 if year in (2003, 2004) else 15
    if 1987 <= year <= 1997: pct = 'half'
    elif year in (2003, 2004): pct = 'wins'
    elif year in (2009, 2010): pct = 'tieLoss'
    else: pct = 'wl'
    out = {
        'maxInnings': inn, 'maxInningsSource': 'disputed' if year == 2003 else 'confirmed',
        'postseasonMaxInnings': post, 'postseasonMaxInningsSource': 'disputed' if year in (2003, 2008) else 'confirmed',
        'winPct': pct, 'winPctSource': 'confirmed',
    }
    if year == 2020:
        out['nineInningsDhMonday'] = 1
    if year == 2021:
        out['nineInningsFrom'] = 810
    return out


# ---- 수상·구장·올스타전 (awards.json, parks.json, allstar_history.json)
# 나무위키 덤프에서 뽑은 가공 CSV (tools/wiki/s6_awards.py, data-src/구장.csv·구장_사용.csv). 2026-10-23 빈 자료 결정 4.
# 게임 코드는 아직 읽지 않는다 (구장 효과·수상 표시는 따로). 선수 구성에 닿지 않아 저장 호환은 그대로.

def read_src_csv(path):
    if not path.exists():
        return []
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def num(v):
    """CSV 칸 → 정수·소수, 빈 칸은 None"""
    if v in (None, ''):
        return None
    f = float(v)
    return int(f) if f.is_integer() else f


def build_history(out_dir):
    src = ROOT / 'data-src'
    awards = [{'award': r['상'], 'year': int(r['연도']), 'part': r['부문'], 'name': r['이름'], 'id': r['선수ID'] or None,
               'team': r['구단'], 'franchise': r['계보'] or None} for r in read_src_csv(src / '수상.csv')]
    (out_dir / 'awards.json').write_text(json.dumps({
        'version': 1, 'source': 'data-src/수상.csv (나무위키 덤프 2026-08-29, tools/wiki/s6_awards.py)',
        'note': 'id가 null이면 선수 ID와 짝짓지 못한 수상 (reports/wiki-awards.md)', 'awards': awards,
    }, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    specs = [{
        'id': r['구장ID'], 'name': r['이름'], 'city': r['도시'], 'from': int(r['시작']), 'to': num(r['끝']),
        'specChange': r['제원변경'] == '1', 'franchise': r['계보'] or None, 'seats': num(r['좌석']),
        'fence': {k: num(r[c]) for k, c in (('l', '좌'), ('lc', '좌중'), ('c', '중'), ('rc', '우중'), ('r', '우'))},
        'height': {'lr': num(r['높이_좌우']), 'c': num(r['높이_중앙'])},
        'turf': {'infield': r['잔디_내야'] or None, 'outfield': r['잔디_외야'] or None}, 'dome': r['돔'] == '1',
        'source': r['출처'], 'confirmed': r['확인'], 'note': r['비고'],
    } for r in read_src_csv(src / '구장.csv')]
    use = [{'franchise': r['계보'], 'park': r['구장ID'], 'from': int(r['시작']), 'to': num(r['끝']), 'kind': r['구분'],
            'source': r['출처'], 'confirmed': r['확인'], 'note': r['비고']} for r in read_src_csv(src / '구장_사용.csv')]
    (out_dir / 'parks.json').write_text(json.dumps({
        'version': 1, 'source': 'data-src/구장.csv · 구장_사용.csv (나무위키 덤프 2026-08-29)',
        'note': 'specs: 구장 × 제원 기간 (to null = 지금까지, franchise 값이 있으면 그 구단 홈 경기에만 쓰는 제원). use: 구단 계보별 주 구장',
        'specs': specs, 'use': use,
    }, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    years = {}
    for r in read_src_csv(src / '올스타.csv'):
        y = years.setdefault(int(r['연도']), {'year': int(r['연도']), 'games': [], 'mvp': None})
        y['games'].append({'n': int(r['경기']), 'date': r['날짜'], 'park': r['구장'], 'parkId': r['구장ID'] or None,
                           'teams': [r['팀1'], r['팀2']], 'score': [int(r['점수1']), int(r['점수2'])], 'win': r['승리']})
    for a in awards:
        if a['award'] == '올스타전 MVP':
            years.setdefault(a['year'], {'year': a['year'], 'games': [], 'mvp': None})['mvp'] = \
                {'name': a['name'], 'id': a['id'], 'team': a['team']}
    (out_dir / 'allstar_history.json').write_text(json.dumps({
        'version': 1, 'source': 'data-src/올스타.csv · 수상.csv 올스타전 MVP (나무위키 덤프 2026-08-29)',
        'years': [years[y] for y in sorted(years)],
    }, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'수상 {len(awards)}줄 (짝 {sum(1 for a in awards if a["id"])}), 구장 제원 {len(specs)}줄 · 사용 {len(use)}줄, '
          f'올스타전 {len(years)}개 해 → {out_dir}')


# ---- 수비 이닝·평판 (defense.json, 능력치 v2 16-3a)
# 가공 CSV: data-src/수비이닝.csv · 수비평판.csv (tools/build_defense.py가 저장소 밖 원본에서 만든다).
# 게임 코드는 아직 읽지 않는다 (판정 함수 src/engine/fielding.ts만, 16-3에서 엔진에 잇는다). 선수 구성에 닿지 않아 저장 호환은 그대로.

DEF_POS = ('C', '1B', '2B', '3B', 'SS', 'LF', 'CF', 'RF')
DEF_FIELDS = ('G', 'GS', '이닝x3', 'E', 'PO', 'A', 'DP', 'SB', 'CS', 'PB')


def build_defense(out_dir):
    """선수 ID → 연도 → {p: 자리 → [G, GS, 이닝×3, E, PO, A, DP, SB, CS, PB], pa: 타석, nd: 수비 없는 타석 (근사)},
    평판: 선수 ID → [[자리, 수준 −2~2, 시기 시작, 시기 끝]] (시기가 없으면 null = 통산).
    수비 없는 타석 = 시즌 타석 × 수비 없는 경기 ÷ 시즌 경기, 수비 없는 경기 = 시즌 경기 − 수비 출장 경기 합 (구단 계보별, 0 아래는 0).
    조사 직원 정리본(선수시즌_수비요약.csv)의 '수비없는_타석_근사'와 같은 식이다. 수비 기록이 시작된 해(2001)부터
    그 해 수비 행이 없는 타자는 타석 전부를 수비 없는 타석으로 센다 (지명타자·대타만 나온 시즌)."""
    src = ROOT / 'data-src'
    rows = read_src_csv(src / '수비이닝.csv')
    if not rows:
        print('수비 이닝 CSV 없음, defense.json을 만들지 않음 (tools/build_defense.py)')
        return
    meta = json.loads((out_dir / 'meta.json').read_text(encoding='utf-8'))
    fr_of = {(s['name'], y): fr for fr, spans in meta['franchises'].items() for s in spans for y in range(s['from'], s['to'] + 1)}
    players = defaultdict(dict)
    def_g = defaultdict(int)  # (선수, 연도, 계보) → 수비 출장 경기 합
    for r in rows:
        pid, y = r['선수ID'], int(r['연도'])
        yr = players[pid].setdefault(str(y), {'p': {}, 'pa': 0, 'nd': 0})
        cur = yr['p'].setdefault(r['자리'], [0] * len(DEF_FIELDS))
        for i, k in enumerate(DEF_FIELDS):
            cur[i] += int(r[k])
        def_g[(pid, y, fr_of.get((r['구단'], y), r['구단']))] += int(r['G'])
    first = min(int(r['연도']) for r in rows)
    last = max(int(r['연도']) for r in rows)
    for y in meta['years']:
        if not first <= y <= last:
            continue
        s = json.loads((out_dir / 'seasons' / f'{y}.json').read_text(encoding='utf-8'))
        f = s['batFields']
        ii, it, ig, ip = f.index('id'), f.index('team'), f.index('g'), f.index('pa')
        for b in s['bat']:
            pid, g, pa = b[ii], b[ig] or 0, b[ip] or 0
            if pa <= 0:
                continue
            yr = players[pid].setdefault(str(y), {'p': {}, 'pa': 0, 'nd': 0})
            dg = def_g.get((pid, y, fr_of.get((b[it], y), b[it])), 0)
            nd_g = max(0, g - dg)
            yr['pa'] += pa
            yr['nd'] += round(pa * nd_g / g) if g > 0 else 0
    for pid in players:
        for yr in players[pid].values():
            if not yr['p']:
                del yr['p']
    rep = defaultdict(list)
    arm = defaultdict(list)
    for r in read_src_csv(src / '수비평판.csv'):
        rep[r['선수ID']].append([r['자리'], int(r['수준']), num(r['시기_시작']), num(r['시기_끝'])])
        if int(r.get('어깨') or 0):
            arm[r['선수ID']].append([int(r['어깨']), num(r['시기_시작']), num(r['시기_끝'])])
    (out_dir / 'defense.json').write_text(json.dumps({
        'version': 1,
        'source': 'data-src/수비이닝.csv (KBO 공식 사이트 수비 기록, 사용자 수집, 가공) · 수비평판.csv (나무위키·위키백과 서술 요약, 등급·시기만)',
        'fields': ['g', 'gs', 'outs', 'e', 'po', 'a', 'dp', 'sb', 'cs', 'pb'], 'from': first, 'to': last,
        'note': 'players: 선수 ID → 연도 → {p: 자리 → fields 순 값 (outs = 이닝 × 3), pa: 타석, nd: 수비 없는 타석 근사}. '
                'rep: 선수 ID → [자리(C…RF·IF·OF·전반), 수준 −2(매우 나쁨)~2(매우 좋음), 시기 시작, 시기 끝] (null = 통산). '
                'arm: 선수 ID → [어깨 1 좋음 / −1 나쁨, 시기 시작, 시기 끝] (평판 글의 어깨 낱말, 능력치 v2 12.6. 아직 엔진에 쓰지 않음)',
        'players': {k: players[k] for k in sorted(players)}, 'rep': {k: rep[k] for k in sorted(rep)}, 'arm': {k: arm[k] for k in sorted(arm)},
    }, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'수비 이닝 {len(rows):,}행 → 선수 {len(players):,}명 ({first}~{last}), 평판 {sum(len(v) for v in rep.values()):,}행 ({len(rep):,}명) → {out_dir / "defense.json"}')


# ---- 구종 (pitches.json, 능력치 v2 16-1a·16-1b)
# 가공 CSV: data-src/구종.csv (visualbaseball.com) · 위키_구종.csv (나무위키·위키백과 덤프). tools/build_pitches.py가 저장소 밖 원본에서 만든다.
# 게임 코드는 아직 읽지 않는다 (16-1·16-2 짐작 식, 16-6 스킬, 16-7 화면에서 쓴다). 선수 구성에 닿지 않아 저장 호환은 그대로.
# prior는 기획안 15.5의 측정 표를 이 CSV에서 다시 낸다 (손으로 옮기지 않는다): 구종 수별 순위 몫, 연도별 포심 평균 구속,
# 같은 투수·시즌 포심과의 구속 차이, 구종 가치(RV/100) 등급 표.

PITCH_FIELDS = ('구속', '수직무브', '수평무브', '릴리스높이', '낙차', '사용률', '투구수', 'RV100')
RV_MIN_N = 300        # 구종 가치 등급 표의 기준 행: 투구수 300 이상 (15.5 ④ 가 2)
RV_MIN_YEAR_ROWS = 30  # 그 해 기준 행이 이만큼 있어야 그 해 중앙값, 아니면 전체 중앙값 (임시값, 15.5 ④ 가 3)
RV_OWN_TABLE = 100     # 기준 행이 이만큼 있어야 자기 표, 아니면 RV_BORROW (15.5 ④ 가 4)
RV_BORROW = {'스위퍼': '슬라이더', '싱커': '투심'}


def _median(v):
    v = sorted(v)
    n = len(v)
    return (v[n // 2] + v[(n - 1) // 2]) / 2


def pitch_prior(rows):
    """15.5의 측정 표. rows = data-src/구종.csv 행"""
    seasons = defaultdict(dict)  # (선수, 연도) → 구종 → 행
    for r in rows:
        seasons[(r['선수ID'], int(r['연도']))][r['구종']] = r
    # ② 구종 수 k별 순위 몫 (사용률이 있는 투수·시즌, 사용률 순, 합 100으로 맞춤)
    by_k = defaultdict(list)
    for d in seasons.values():
        u = sorted((float(r['사용률']) for r in d.values() if r['사용률'] not in ('', None)), reverse=True)
        if not u or sum(u) <= 0:
            continue
        t = sum(u)
        by_k[len(u)].append([100 * x / t for x in u])
    share = {str(k): {'n': len(v), 's': [round(sum(x[i] for x in v) / len(v), 1) for i in range(k)]} for k, v in sorted(by_k.items())}
    # ③ 연도별 포심 평균 구속 (투수·시즌 단순 평균), 같은 투수·시즌 포심과의 차이 [평균, 표준편차, 행 수]
    fb = defaultdict(list)
    diff = defaultdict(list)
    for (pid, y), d in seasons.items():
        f = d.get('포심')
        if not f or f['구속'] in ('', None):
            continue
        v0 = float(f['구속'])
        fb[y].append(v0)
        for p, r in d.items():
            if p != '포심' and r['구속'] not in ('', None):
                diff[p].append(float(r['구속']) - v0)
    fb_speed = {str(y): round(sum(v) / len(v), 1) for y, v in sorted(fb.items())}
    speed_diff = {}
    for p, v in diff.items():
        m = sum(v) / len(v)
        sd = math.sqrt(sum((x - m) ** 2 for x in v) / (len(v) - 1)) if len(v) > 1 else 0
        speed_diff[p] = [round(m, 1), round(sd, 1), len(v)]
    # ④ 가 구종 가치 등급 표: 기준 행(투구수 300 이상)에서 c = RV − m(그 해, 구종), 구종별로 정렬한 c
    base = defaultdict(list)  # 구종 → [(연도, RV)]
    for r in rows:
        if r['투구수'] not in ('', None) and r['RV100'] not in ('', None) and int(r['투구수']) >= RV_MIN_N:
            base[r['구종']].append((int(r['연도']), float(r['RV100'])))
    tables = {}
    for p, v in base.items():
        if len(v) < RV_OWN_TABLE:
            continue
        by_y = defaultdict(list)
        for y, x in v:
            by_y[y].append(x)
        med = {str(y): round(_median(xs), 3) for y, xs in sorted(by_y.items()) if len(xs) >= RV_MIN_YEAR_ROWS}
        med['all'] = round(_median([x for _, x in v]), 3)
        c = sorted(round(x - med.get(str(y), med['all']), 3) for y, x in v)
        tables[p] = {'med': med, 'c': c}
    return {
        'share': share, 'fbSpeed': fb_speed, 'speedDiff': speed_diff,
        'rv': {'minN': RV_MIN_N, 'baseRows': {p: len(v) for p, v in sorted(base.items(), key=lambda x: -len(x[1]))},
               'borrow': RV_BORROW, 'tables': tables},
    }


def build_pitches(out_dir):
    """선수 ID → 연도 → 구종 → [구속, 수직 무브, 수평 무브, 릴리스 높이, 낙차, 사용률, 투구수, RV/100] (빈 칸 null),
    prior (15.5 측정 표), wiki: names(원문 낱말 목록) + 선수 ID → {p: [[표준 구종, 원문 색인, 주무기 0/1, 시작, 끝]] (위키 순서), top: 최고 구속}"""
    src = ROOT / 'data-src'
    rows = read_src_csv(src / '구종.csv')
    if not rows:
        print('구종 CSV 없음, pitches.json을 만들지 않음 (tools/build_pitches.py)')
        return
    players = defaultdict(lambda: defaultdict(dict))
    for r in rows:
        players[r['선수ID']][r['연도']][r['구종']] = [num(r[k]) for k in PITCH_FIELDS]
    wiki_rows = read_src_csv(src / '위키_구종.csv')
    names, name_at = [], {}
    wiki = defaultdict(lambda: {'p': [], 'top': None})
    for r in sorted(wiki_rows, key=lambda r: (r['선수ID'], int(r['순서']))):
        raw = r['구종_원문']
        if raw not in name_at:
            name_at[raw] = len(names)
            names.append(raw)
        w = wiki[r['선수ID']]
        w['p'].append([r['구종'], name_at[raw], 1 if r['주무기'] == '1' else 0, num(r['시작연도']), num(r['끝연도'])])
        top = num(r['최고구속'])
        if top is not None and (w['top'] is None or top > w['top']):
            w['top'] = top
    years = sorted({int(r['연도']) for r in rows})
    (out_dir / 'pitches.json').write_text(json.dumps({
        'version': 1,
        'source': 'data-src/구종.csv (visualbaseball.com 구종 자료, 사용자 수집 2026-10, 가공) · 위키_구종.csv (나무위키·위키백과 덤프, 구종 낱말·순서·주무기·시기·최고 구속만)',
        'fields': ['speed', 'vmov', 'hmov', 'rel', 'drop', 'usage', 'n', 'rv'], 'from': years[0], 'to': years[-1],
        'note': 'players: 선수 ID → 연도 → 구종 → fields 순 값 (구속 km/h, 무브·릴리스 cm, 낙차 °, 사용률 %, 투구수, RV/100 — 음수가 투수에게 좋음, 화면에 내지 않는다). '
                'wiki.players: 선수 ID → {p: [표준 구종, 원문 색인(wiki.names), 주무기, 시작, 끝(null = 커리어 전체)], top: 최고 구속}',
        'players': {k: players[k] for k in sorted(players)},
        'prior': pitch_prior(rows),
        'wiki': {'names': names, 'players': {k: wiki[k] for k in sorted(wiki)}},
    }, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'구종 {len(rows):,}행 → 선수 {len(players):,}명 ({years[0]}~{years[-1]}), 위키 구종 {len(wiki_rows):,}행 ({len(wiki):,}명) → {out_dir / "pitches.json"}')


# ---- 규칙 변경 (rule_changes.json, 기획안 docs/리그규칙변경_기획안.md 2.3)

RULE_CAT = {'공인구': 'ball', '경기': 'game', '엔트리': 'roster', '외국인': 'foreign', 'FA': 'fa', '드래프트': 'draft', '샐러리캡': 'cap', '리그': 'league'}
RULE_EFFECT = {'표시만': 'show', '표시만(평균)': 'show-env', '리그 운영': 'league', '경기': 'engine'}
RULE_STATUS = {'자료집 확인': 'confirmed', '조사값': 'research', '출처 갈림': 'conflict', '미확인': 'unknown'}
# 키가 있는 항목은 전·후 값을 게임이 실제로 쓰는 값에서 읽는다 (src/league/ruleNews.ts의 RULE_KEYS와 같아야 한다)
RULE_KEYS = {'rosterSize', 'foreignLimit', 'maxInnings', 'postseasonTeams', 'faEnabled', 'faService', 'salaryCap'}


def build_rule_changes(out_dir):
    """규칙변경.csv + 공인구.csv → rule_changes.json. 선수 월드 지문(선수 ID·소속)에 들어가지 않으므로 고쳐도 예전 저장을 쓴다."""
    if not RULES_CSV.exists():
        return
    items = []
    with RULES_CSV.open(encoding='utf-8-sig', newline='') as f:
        for n, r in enumerate(csv.DictReader(f), start=2):
            where = f'규칙변경.csv {n}행'
            def pick(table, col):
                v = (r[col] or '').strip()
                if v not in table:
                    raise SystemExit(f'{where}: {col} 값 "{v}"는 {"/".join(table)} 중 하나여야 한다')
                return table[v]
            it = {'year': int(r['연도'])}
            if (r['날짜'] or '').strip():
                it['date'] = r['날짜'].strip()
            it['cat'] = pick(RULE_CAT, '분류')
            if (r['키'] or '').strip():
                it['key'] = r['키'].strip()
            it['title'] = r['제목'].strip()
            for col, k in (('전', 'before'), ('후', 'after'), ('단위', 'unit')):
                if (r[col] or '').strip():
                    it[k] = r[col].strip()
            it['effect'] = pick(RULE_EFFECT, '반영')
            it['status'] = pick(RULE_STATUS, '확인')
            if (r['출처'] or '').strip():
                it['source'] = r['출처'].strip()
            note = (r['비고'] or '').strip()
            if note.startswith('게임은 단순화:'):
                it['simplified'] = note[len('게임은 단순화:'):].strip()
            elif note:
                it['note'] = note
            items.append(it)
    balls = defaultdict(dict)
    if BALL_CSV.exists():
        with BALL_CSV.open(encoding='utf-8-sig', newline='') as f:
            for r in csv.DictReader(f):
                y = int(r['연도'])
                b = balls[y]
                b['year'] = y
                kind, val, note = r['구분'].strip(), (r['값'] or '').strip(), (r['비고'] or '').strip()
                status = RULE_STATUS[r['확인'].strip()]
                if kind == '기준':
                    lo, hi = val.split('~')
                    b['range'] = [float(lo), float(hi)]
                    if note:
                        b['note'] = note
                elif kind == '단일':
                    b['single'] = val == '단일'
                elif kind == '검사':
                    t = {'round': int(r['차수'])}
                    if (r['날짜'] or '').strip():
                        t['date'] = r['날짜'].strip()
                    if val:
                        t['avg'] = float(val)
                    if note:
                        t['note'] = note
                    b.setdefault('tests', []).append(t)
                else:
                    raise SystemExit(f'공인구.csv: 구분 "{kind}"는 기준/단일/검사 중 하나여야 한다')
                b['status'] = status
                if (r['출처'] or '').strip():
                    b['source'] = r['출처'].strip()
    (out_dir / 'rule_changes.json').write_text(json.dumps({
        'version': 1,
        'note': '연도별 리그 규칙 변경 (data-src/규칙변경.csv, 공인구.csv). status research = 웹 조사 값, 자료집 제도연표 확인 전. '
                'key가 있는 항목의 전·후 값은 게임 규칙 값에서 계산한다.',
        'items': items,
        'ball': [balls[y] for y in sorted(balls)],
    }, ensure_ascii=False, indent=1), encoding='utf-8')
    print(f'규칙 변경: {len(items)}건, 공인구 {len(balls)}개 연도 → {out_dir / "rule_changes.json"}')


# ---- 감독 (managers.json)

# 양대 리그 시대(1999·2000)는 리그마다 최종 순위 1위가 있어 한국시리즈 우승 구단(그 해 이름)을 따로 둔다
# (src/league/history.ts TWO_LEAGUE_FINAL과 같은 값)
TWO_LEAGUE_CHAMPION = {1999: '한화', 2000: '현대'}

# 문자중계 팀 코드 → 계보 (data-src/pbp_derived/team_pitching.csv)
PBP_TEAM = {'HH': '한화', 'HT': 'KIA', 'KT': 'KT', 'LG': 'LG', 'LT': '롯데', 'NC': 'NC', 'OB': '두산', 'SK': 'SSG', 'SS': '삼성', 'WO': '히어로즈'}


def build_managers(out_dir):
    """감독.csv → {seasons: {연도: {계보: [{name, kind, mid}]}}, people: {이름: 경력·운영 경향}}.
    운영 경향은 그 감독이 개막부터 끝까지 맡은 시즌(중도 교체가 없는 해)의 팀 기록을 그 해 리그 평균과 견준 배율이다.
      steal: (도루 + 도루자) / (단타 + 볼넷 + 사구), bunt: 희생번트 / 타석, relief: (등판 − 선발) / 경기
    시즌이 적으면 1 쪽으로 당긴다 (n / (n + 2)). 성격·인물평은 넣지 않는다 (기록 사실만).
    투구 단위 자료가 있는 해(2023~, data-src/pbp_derived/team_pitching.csv, 엔진 2 10단계)는 더 본다:
      push: 팀 선발 평균 투구 수 / 리그, pen: 팀 구원 연투(전날 등판) 비율 / 리그 (감독 성향 '투수 밀어붙이기'·'불펜 연투 허용').
      개막 감독에게는 그 해 팀 값을 byYear에 그대로 남기고 (시즌 중 교체가 있던 해도), 시즌 내내 맡은 해의 평균은 위처럼 당겨 push·pen에.
      자료가 없는 감독은 push·pen이 없다 (엔진은 push 대신 relief, pen은 보통).
    엔진 2 13단계에서 더 본다 (같은 자료, 같은 방식으로 byYear·평균):
      ibb: 팀 고의4구 ÷ 상대 타석을 2023~2025(온전한 시즌) 리그 합친 비율로 나눈 값 (감독 성향 '고의4구 빈도'). 해마다 리그 수준 차이도 담는다
        (2024 리그 0.88배). 그래서 세 해 평균은 1 근처.
      cond: 팀 선발 라인업 변동(전 경기와 바뀐 선발 수 평균, team_lineup.csv) ÷ 그 해 리그 (감독 성향 '컨디션 반영' 대리값)."""
    if not MANAGERS_CSV.exists():
        return
    meta = json.loads((out_dir / 'meta.json').read_text(encoding='utf-8'))
    with MANAGERS_CSV.open(encoding='utf-8-sig', newline='') as f:
        rows_ = [r for r in csv.DictReader(f)]
    seasons = defaultdict(lambda: defaultdict(list))
    for r in rows_:
        seasons[r['시즌']][r['계보']].append({'name': r['감독'], 'kind': r['구분'], **({'mid': True} if r['시점'] == '중도' else {})})

    def team_stats(y):
        sd = json.loads((out_dir / 'seasons' / f'{y}.json').read_text(encoding='utf-8'))
        bi = {k: i for i, k in enumerate(sd['batFields'])}
        pi = {k: i for i, k in enumerate(sd['pitFields'])}
        out = {}
        for t in sd['teams']:
            b = [r for r in sd['bat'] if r[bi['team']] == t['name']]
            p = [r for r in sd['pit'] if r[pi['team']] == t['name']]
            g = lambda rs, k, ix: sum((r[ix[k]] or 0) for r in rs)
            singles = g(b, 'h', bi) - g(b, 'd', bi) - g(b, 't', bi) - g(b, 'hr', bi)
            on1 = singles + g(b, 'bb', bi) + g(b, 'hbp', bi)
            out[t['franchise']] = {
                'steal': (g(b, 'sb', bi) + g(b, 'cs', bi)), 'on1': on1, 'sac': g(b, 'sac', bi), 'pa': g(b, 'pa', bi),
                'relief': g(p, 'g', pi) - g(p, 'gs', pi), 'games': t['g'], 'rank': t['rank'], 'name': t['name'],
            }
        return out

    acc = defaultdict(lambda: {'steal': [], 'bunt': [], 'relief': []})
    people = {}
    for y in meta['years']:
        ts = team_stats(y)
        lg = {k: sum(v[k] for v in ts.values()) for k in ('steal', 'on1', 'sac', 'pa', 'relief', 'games')}
        for fr, ms in seasons.get(str(y), {}).items():
            for m in ms:
                pe = people.setdefault(m['name'], {'first': y, 'last': y, 'seasons': 0, 'acting': 0, 'teams': [], 'titles': 0})
                pe['last'] = y
                if fr not in pe['teams']:
                    pe['teams'].append(fr)
                if m['kind'] == '정식':
                    pe['seasons'] += 1
                else:
                    pe['acting'] += 1
            full = [m for m in ms if not m.get('mid')]
            if len(ms) != 1 or not full or fr not in ts:
                continue
            name = full[0]['name']
            t = ts[fr]
            champion = (t['name'] == TWO_LEAGUE_CHAMPION[y]) if y in TWO_LEAGUE_CHAMPION else t['rank'] == 1
            if champion and y not in meta.get('inProgress', []):  # 진행 중인 시즌의 1위는 우승이 아니다
                people[name]['titles'] += 1
            ratio = lambda a, b, la, lb: (a / b) / (la / lb) if b and lb and la else None
            for k, v in (('steal', ratio(t['steal'], t['on1'], lg['steal'], lg['on1'])), ('bunt', ratio(t['sac'], t['pa'], lg['sac'], lg['pa'])),
                         ('relief', ratio(t['relief'], t['games'], lg['relief'], lg['games']))):
                if v is not None:
                    acc[name][k].append(v)
    # 투구 단위 자료 (2023~): 구단-시즌 → push·pen 배율
    pitch = defaultdict(dict)
    tp = PBP_DERIVED / 'team_pitching.csv'
    if tp.exists():
        with tp.open(encoding='utf-8', newline='') as f:
            trows = list(csv.DictReader(f))
        for y in sorted({r['year'] for r in trows}):
            ys = [r for r in trows if r['year'] == y]
            starts = sum(int(r['starts']) for r in ys)
            lg_mean = sum(int(r['starts']) * float(r['starter_mean']) for r in ys) / starts
            apps = sum(int(r['relief_apps']) for r in ys)
            lg_b2b = sum(int(r['relief_apps']) * float(r['back_to_back_rate']) for r in ys) / apps
            for r in ys:
                pitch[y][PBP_TEAM[r['team']]] = {'push': round(float(r['starter_mean']) / lg_mean, 3),
                                                 'pen': round(float(r['back_to_back_rate']) / lg_b2b, 3)}
        # 고의4구: 온전한 시즌을 합친 리그 비율이 기준 (해마다 리그 수준 차이를 남긴다)
        full = [r for r in trows if r['partial'] == '0']
        ibb_base = sum(int(r['ibb']) for r in full) / sum(int(r['bf']) for r in full)
        for r in trows:
            pitch[r['year']][PBP_TEAM[r['team']]]['ibb'] = round(int(r['ibb']) / int(r['bf']) / ibb_base, 3)
    tl = PBP_DERIVED / 'team_lineup.csv'
    if tl.exists():
        with tl.open(encoding='utf-8', newline='') as f:
            lrows = list(csv.DictReader(f))
        for y in sorted({r['year'] for r in lrows}):
            ys = [r for r in lrows if r['year'] == y]
            lg_moved = sum(int(r['games']) * float(r['mean_changed']) for r in ys) / sum(int(r['games']) for r in ys)
            for r in ys:
                pitch[y].setdefault(PBP_TEAM[r['team']], {})['cond'] = round(float(r['mean_changed']) / lg_moved, 3)
    pacc = defaultdict(lambda: {'push': [], 'pen': [], 'ibb': [], 'cond': []})
    for y, teams in pitch.items():
        for fr, v in teams.items():
            ms = seasons.get(y, {}).get(fr, [])
            opening = [m for m in ms if not m.get('mid')]
            if not opening or opening[0]['name'] not in people:
                continue
            name = opening[0]['name']
            people[name].setdefault('byYear', {})[y] = v
            if len(ms) == 1 and y not in [str(x) for x in meta.get('inProgress', [])]:
                for k in ('push', 'pen', 'ibb', 'cond'):
                    if k in v:
                        pacc[name][k].append(v[k])
    for name, a in pacc.items():
        for k in ('push', 'pen', 'ibb', 'cond'):
            vs = a[k]
            if vs:
                people[name][k] = round(1 + (sum(vs) / len(vs) - 1) * len(vs) / (len(vs) + 2), 3)
    for name, pe in people.items():
        a = acc.get(name)
        n = len(a['steal']) if a else 0
        pe['full'] = n
        for k in ('steal', 'bunt', 'relief'):
            vs = a[k] if a else []
            mean = sum(vs) / len(vs) if vs else 1.0
            pe[k] = round(1 + (mean - 1) * len(vs) / (len(vs) + 2), 3)
    (out_dir / 'managers.json').write_text(json.dumps({
        'source': MANAGERS_CSV.name,
        'seasons': seasons,
        'people': people,
    }, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'감독: {len(rows_)}행, {len(people)}명 → {out_dir / "managers.json"}')


# ---- 병역 (military.json)

SERVICE_UNITS = ('상무', '경찰')


def service_unit(career):
    """2026 프로필 경력에 상무·경찰 야구단이 있으면 그 이름. 학교 이름(군산상고 등)과 섞이지 않게 '-'로 나눈 칸 단위로 본다"""
    for part in (career or '').split('-'):
        part = part.strip()
        for u in SERVICE_UNITS:
            if part.startswith(u):
                return u
    return None


def build_ppa():
    """선수별 타석당 투구 수 → public/data/ppa.json (경기 엔진 2 D단계, 사슬 길이 손잡이).
    {bat|pit: {선수 ID: {연도: [타석, 리그 대비 비율]}}}. 비율 = 그 선수 P/PA ÷ 같은 자료·같은 정의의 그 해 리그 P/PA.
    투구 단위 자료(2023~, data-src/pbp_derived/{batter,pitcher}_ppa.csv, 고의4구 뺌)가 먼저, 투수 2015~2022는 하드힛 월별 표
    (data-src/hardhit_derived/pitcher_ppa.csv, 자동 고의4구 포함 정의, 리그 값은 league_ppa.csv)."""
    out = {'bat': defaultdict(dict), 'pit': defaultdict(dict)}
    lg = {}
    with (PBP_DERIVED / 'pa_pitches.csv').open(encoding='utf-8') as f:
        lg = {int(r['year']): float(r['mean']) for r in csv.DictReader(f) if r['group'] == 'all'}
    for side, name in (('bat', 'batter_ppa.csv'), ('pit', 'pitcher_ppa.csv')):
        with (PBP_DERIVED / name).open(encoding='utf-8') as f:
            for r in csv.DictReader(f):
                y = int(r['year'])
                out[side][r['id']][str(y)] = [int(r['pa']), round(float(r['ppa']) / lg[y], 4)]
    with (HARDHIT_DERIVED / 'league_ppa.csv').open(encoding='utf-8') as f:
        hlg = {int(r['year']): float(r['ppa']) for r in csv.DictReader(f)}
    acc = defaultdict(lambda: [0, 0])
    with (HARDHIT_DERIVED / 'pitcher_ppa.csv').open(encoding='utf-8') as f:
        for r in csv.DictReader(f):
            if int(r['year']) < min(lg):
                a = acc[(r['id'], int(r['year']))]
                a[0] += int(r['pa'])
                a[1] += int(r['pitches'])
    for (pid, y), (pa, pitches) in sorted(acc.items()):
        if pa > 0 and pitches > 0:
            out['pit'][pid][str(y)] = [pa, round(pitches / pa / hlg[y], 4)]
    (OUT / 'ppa.json').write_text(json.dumps({
        'version': 1,
        'note': '선수별 타석당 투구 수의 리그 대비 비율 [타석, 비율]. 투구 단위 자료(slothman3878 KBO Play-by-Play v0, CC BY 4.0) 2023~, '
                '투수 2015~2022는 하드힛(https://hardhit.ai) 월별 표. tools/build_data.py build_ppa',
        'bat': out['bat'], 'pit': out['pit'],
    }, ensure_ascii=False, separators=(',', ':'), sort_keys=True), encoding='utf-8')
    return sum(len(v) for v in out['bat'].values()), sum(len(v) for v in out['pit'].values())


def build_running():
    """선수별 추가 진루·번트 → public/data/running.json (능력치 v2 16-3d, 주루 판단·번트 수행 실측 시작값).
    {bat: {선수 ID: {연도: [기회, 성공, 주루사, 리그 기대 성공, 희생번트, W·V 공]}}, lg: {연도: [희생번트, W·V 공]}}.
    가공 CSV는 tools/pbp_running.py (data-src/pbp_derived/runner_adv.csv·batter_bunt.csv, 투구 단위 자료 2023~)."""
    bat = defaultdict(dict)
    lg = defaultdict(lambda: [0, 0])
    with (PBP_DERIVED / 'runner_adv.csv').open(encoding='utf-8') as f:
        for r in csv.DictReader(f):
            bat[r['id']][r['year']] = [int(r['opp']), int(r['ext']), int(r['outs']), float(r['exp']), 0, 0]
    with (PBP_DERIVED / 'batter_bunt.csv').open(encoding='utf-8') as f:
        for r in csv.DictReader(f):
            row = bat[r['id']].setdefault(r['year'], [0, 0, 0, 0.0, 0, 0])
            row[4], row[5] = int(r['sh']), int(r['wv'])
            lg[r['year']][0] += int(r['sh'])
            lg[r['year']][1] += int(r['wv'])
    (OUT / 'running.json').write_text(json.dumps({
        'version': 1,
        'note': '선수별 추가 진루 [기회, 성공, 주루사, 리그 기대 성공]과 번트 [희생번트, 번트 파울·헛스윙 공]. 투구 단위 자료(slothman3878 '
                'KBO Play-by-Play v0, CC BY 4.0) 2023~, 2026은 7월 26일까지. tools/pbp_running.py → tools/build_data.py build_running',
        'bat': bat, 'lg': lg,
    }, ensure_ascii=False, separators=(',', ':'), sort_keys=True), encoding='utf-8')
    return len(bat)


def build_plate():
    """선수별 볼카운트 사슬 실측 → public/data/plate.json (능력치 v2 16-1, 엔진 src/engine/plate.ts·ratings.ts plateResidual).
    {bat|pit: {선수 ID: {연도: [건수...]}}}. 타자 [존 안 공, 존 안 스윙, 기대, 존 밖 공, 존 밖 스윙, 기대, 스윙, 맞힘, 기대,
    2스트라이크 스윙, 맞힘, 기대, 맞힘, 파울, 기대], 투수 [공, 존 안 공, 기대, 스윙, 맞힘, 기대, 맞힘, 파울, 기대].
    기대 = 같은 칸·존에서 리그 평균이 냈을 값. 투구 단위 자료(2023~, data-src/pbp_derived/{batter,pitcher}_plate.csv, tools/pbp_plate.py)."""
    cols = {
        'bat': ('batter_plate.csv', ['n_in', 'sw_in', 'exp_sw_in', 'n_out', 'sw_out', 'exp_sw_out', 'swings', 'contact', 'exp_contact',
                                     'swings_2s', 'contact_2s', 'exp_contact_2s', 'contacts', 'fouls', 'exp_fouls']),
        'pit': ('pitcher_plate.csv', ['pitches', 'n_in', 'exp_in', 'swings', 'contact', 'exp_contact', 'contacts', 'fouls', 'exp_fouls']),
    }
    out = {'bat': defaultdict(dict), 'pit': defaultdict(dict)}
    for side, (name, keys) in cols.items():
        with (PBP_DERIVED / name).open(encoding='utf-8') as f:
            for r in csv.DictReader(f):
                out[side][r['id']][r['year']] = [num(r[k]) for k in keys]
    (OUT / 'plate.json').write_text(json.dumps({
        'version': 1,
        'note': '선수별 볼카운트 사슬 실측 건수와 리그 기대값. 투구 단위 자료(slothman3878 KBO Play-by-Play v0, CC BY 4.0) 2023~. '
                'tools/pbp_plate.py → tools/build_data.py build_plate',
        'bat': out['bat'], 'pit': out['pit'],
    }, ensure_ascii=False, separators=(',', ':'), sort_keys=True), encoding='utf-8')
    return sum(len(v) for v in out['bat'].values()), sum(len(v) for v in out['pit'].values())


def build_homes():
    """data-src/연고지.csv → public/data/homes.json: 계보별 연도 구간의 연고 도시와 위·경도 (원정 거리 컨디션, 경기 엔진 2 12단계).
    끝 연도가 빈 칸이면 지금까지. 거리는 엔진(src/league/distance.ts)이 위·경도로 계산한다 (직선 거리, 임시값)."""
    rows = []
    with HOMES_CSV.open(encoding='utf-8-sig') as f:
        for r in csv.DictReader(f):
            rows.append({'franchise': r['계보'], 'from': int(r['시작']), **({'to': int(r['끝'])} if r['끝'] else {}),
                         'city': r['도시'], 'lat': float(r['위도']), 'lon': float(r['경도'])})
    (OUT / 'homes.json').write_text(json.dumps({
        'version': 1,
        'note': '구단 계보별 연고 도시와 위·경도. 출처·확인 상태는 data-src/연고지.csv (모두 임시값).',
        'homes': rows,
    }, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    return rows


SCHEDULE_CSV = ROOT / 'data-src' / '일정규칙.csv'
SCHEDULE_FRAMES = ('W82', 'W83', 'T3W2', 'S2', 'S33', 'S33-99', 'S33-9', 'S33L2', 'S33F')
SCHEDULE_POSTPONE = ('R1', 'R2', 'R3', 'R3c', 'R4', 'R5')
SCHEDULE_STATUS = ('확인', '조사값', '출처 갈림', '미확인', '임시값')
WEEKDAYS = '일월화수목금토'  # JS Date.getDay() 순서


def build_schedule():
    """data-src/일정규칙.csv → public/data/schedule.json: 연도별 정규시즌 일정 규칙 (일정 개편 기획안 2장).
    한 해 한 줄 + '2027~' 줄(자료 없는 해, 날짜를 규칙으로 계산). 날짜는 'MM-DD'.
    '2027~' 줄은 개막·올스타를 {weekday, from, to}(그 구간의 그 요일)로, 올스타 휴식을 올스타 날 기준 [앞, 뒤] 날 수로 둔다.
    선수 구성 지문과 무관하다 (저장 호환 그대로)."""
    def date(s, where):
        if not re.fullmatch(r'\d\d-\d\d', s):
            raise SystemExit(f'일정규칙.csv {where}: 날짜 형식 "{s}" (MM-DD)')
        return s

    def day_rule(s, where):
        m = re.fullmatch(r'(.) (\d\d-\d\d)~(\d\d-\d\d)', s)
        if not m or m.group(1) not in WEEKDAYS:
            raise SystemExit(f'일정규칙.csv {where}: 규칙 날짜 "{s}" (예: 토 03-22~03-28)')
        return {'weekday': WEEKDAYS.index(m.group(1)), 'from': m.group(2), 'to': m.group(3)}

    def ranges(s, where):
        out = []
        for part in filter(None, (x.strip() for x in (s or '').split(';'))):
            m = re.fullmatch(r'(\d\d-\d\d)~(\d\d-\d\d)', part)
            if not m:
                raise SystemExit(f'일정규칙.csv {where}: 기간 형식 "{part}" (MM-DD~MM-DD)')
            out.append({'from': m.group(1), 'to': m.group(2)})
        return out

    years, future = {}, None
    with SCHEDULE_CSV.open(encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(f):
            key = r['연도'].strip()
            is_future = key.endswith('~')
            where = key
            vs = r['상대별'].split('·')
            row = {}
            if r['구단수']:
                row['teams'] = int(r['구단수'])
            row['games'] = int(r['팀당경기'])
            row['vs'] = int(vs[0])
            if len(vs) > 1:
                row['vsOther'] = int(vs[1])  # 다른 리그 상대 (양대 리그 해)
            if r['연전틀'] not in SCHEDULE_FRAMES:
                raise SystemExit(f'일정규칙.csv {where}: 모르는 연전틀 {r["연전틀"]}')
            row['frame'] = r['연전틀']
            row['opening'] = day_rule(r['개막일'], where) if is_future else date(r['개막일'], where)
            for col, k in (('종료참고', 'endRef'), ('종료참고2', 'endRef2'), ('가을시작', 'postseasonStart'), ('잔여발표', 'remainingFrom')):
                if r[col]:
                    row[k] = date(r[col], where)
            breaks = []
            for part in filter(None, (x.strip() for x in (r['쉬는기간'] or '').split(';'))):
                span, _, reason = part.partition(':')
                m = re.fullmatch(r'올스타([+-]\d+)~올스타([+-]\d+)', span)
                if m:  # 올스타 날 기준 (자료 없는 해)
                    row['allStarBreak'] = [int(m.group(1)), int(m.group(2))]
                    continue
                m = re.fullmatch(r'(\d\d-\d\d)~(\d\d-\d\d)', span)
                if not m or not reason:
                    raise SystemExit(f'일정규칙.csv {where}: 쉬는기간 형식 "{part}" (MM-DD~MM-DD:사유)')
                breaks.append({'from': m.group(1), 'to': m.group(2), 'reason': reason})
            row['breaks'] = breaks
            if r['올스타일']:
                row['allStar'] = day_rule(r['올스타일'], where) if is_future else date(r['올스타일'], where)
            row['allStarGamesReal'] = int(r['실제올스타경기수'] or 0)
            if r['순연'] not in SCHEDULE_POSTPONE:
                raise SystemExit(f'일정규칙.csv {where}: 모르는 순연 코드 {r["순연"]}')
            row['postpone'] = r['순연']
            row['dhBan'] = ranges(r['더블헤더금지'], where)
            row['unscheduled'] = int(r['미편성(팀당)'] or 0)
            if r['2연전시작']:  # '올스타'(올스타 휴식 뒤), 'MM-DD'(그날부터 본 일정 끝까지), 'MM-DD~MM-DD'(그 뒤 다시 3연전, 2020)
                v = r['2연전시작']
                m = re.fullmatch(r'(\d\d-\d\d)(?:~(\d\d-\d\d))?', v)
                if v == '올스타':
                    row['twoGameFrom'] = 'allStar'
                elif m:
                    row['twoGameFrom'] = m.group(1)
                    if m.group(2):
                        row['twoGameTo'] = m.group(2)
                else:
                    raise SystemExit(f'일정규칙.csv {where}: 2연전시작은 "올스타", MM-DD, MM-DD~MM-DD ({v})')
            row['rainMid'] = r['경기중비']
            if r['확인'] not in SCHEDULE_STATUS:
                raise SystemExit(f'일정규칙.csv {where}: 확인 값 "{r["확인"]}" (가능: {", ".join(SCHEDULE_STATUS)})')
            row['status'] = r['확인']
            if is_future:
                row['from'] = int(key[:-1])
                future = row
            else:
                years[key] = row
    (OUT / 'schedule.json').write_text(json.dumps({
        'version': 1,
        'note': '연도별 정규시즌 일정 규칙 (일정 개편 기획안 2장). 출처·확인 상태·비고는 data-src/일정규칙.csv.',
        'years': years,
        'future': future,
    }, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    return years, future


def build_military(profiles, out_dir):
    """실제 병역 자료 → {선수 ID: {served?, exempt?}}.
    profiles: 선수 ID → 2026 프로필 경력 문자열. 시즌 기록과 선수 마스터는 out_dir의 JSON에서 읽는다 (같은 빌드에서 만든 것).
    - served: 프로필 경력에 상무·경찰이 있는 선수. end = 1군 기록이 1~2시즌 빈 마지막 해 (복무를 마친 해 추정, 없으면 생략)
    - exempt: 병역_면제명단.csv의 금메달 대표팀(미필이었으면 면제)과 2026 새 면제자. year = 대회 연도"""
    players = json.loads((out_dir / 'players.json').read_text(encoding='utf-8'))['players']
    meta = json.loads((out_dir / 'meta.json').read_text(encoding='utf-8'))
    seen = defaultdict(dict)  # 선수 ID → {연도: 구단}
    for y in meta['years']:
        sd = json.loads((out_dir / 'seasons' / f'{y}.json').read_text(encoding='utf-8'))
        for r in sd['bat'] + sd['pit']:
            seen[str(r[0])][y] = r[1]
    table = {}
    for p in players:
        if p.get('foreign'):
            continue
        unit = service_unit(profiles.get(p['id']))
        if not unit:
            continue
        ys = sorted(seen.get(p['id'], {}))
        end = None
        for a, b in zip(ys, ys[1:]):
            if 2 <= b - a <= 3:  # 1~2시즌 공백
                end = b - 1
        table[p['id']] = {'served': {'unit': unit, **({'end': end} if end else {})}}
    unmatched = []
    if MILITARY_CSV.exists():
        by_name = defaultdict(list)
        for p in players:
            if not p.get('foreign'):
                by_name[p['name']].append(p)
        with MILITARY_CSV.open(encoding='utf-8-sig', newline='') as f:
            for r in csv.DictReader(f):
                y, name, team = int(r['연도']), r['선수명'].strip(), (r['당시소속'] or '').strip()
                cands = by_name.get(name, [])
                if team in ('아마추어',):
                    pick = [p for p in cands if (p.get('first') or 0) >= y]
                elif team.startswith('해외:'):  # 그 해에는 해외 리그. 다른 해에 뛴 KBO 구단으로 가린다
                    pick = [p for p in cands if team[3:] in seen[p['id']].values()]
                elif team in ('', '해외'):
                    pick = cands
                else:
                    pick = [p for p in cands if any(seen[p['id']].get(x) == team for x in (y - 1, y, y + 1))]
                if len(pick) != 1:
                    unmatched.append(f"{y} {name}({team or '-'}): 후보 {len(pick)}명")
                    continue
                e = table.setdefault(pick[0]['id'], {})
                if 'exempt' not in e or e['exempt']['year'] > y:
                    e['exempt'] = {'year': y, 'event': r['대회']}
    real_n = len(table)
    wiki_n, wiki_skip = wiki_military(table, players, seen)
    (out_dir / 'military.json').write_text(json.dumps({
        'version': 1,
        'note': '실제 병역 자료. served = 2026 프로필 경력의 상무·경찰(end = 1군 공백으로 본 복무 끝 해, 추정), '
                'exempt = 국제대회 금메달 대표팀(병역_면제명단.csv, year = 대회 연도). '
                'src = wiki 는 위 자료가 없는 선수를 위키백과·나무위키 병역 칸으로 채운 것 (served.kind active = 현역, reserve = 보충역). '
                '여기 없는 선수는 자료가 없는 것이다.',
        'players': dict(sorted(table.items())),
    }, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    served = sum(1 for v in table.values() if 'served' in v and not v.get('src'))
    exempt = sum(1 for v in table.values() if 'exempt' in v and not v.get('src'))
    print(f'병역: 상무·경찰 {served}명, 금메달 면제·대표 {exempt}명, 위키로 채움 {wiki_n}명 (읽지 못함 {len(wiki_skip)}) → {out_dir / "military.json"}')
    for u in unmatched:
        print(f'  병역_면제명단.csv 연결 못 함: {u}')
    return {'real': real_n, 'wiki': wiki_n, 'skip': wiki_skip}


# ---- 위키 가공값 (data-src/wiki/, 빈 칸만 채운다)

def wiki_profile():
    """인적사항.csv → 선수 ID → {항목: (값, 출처)}. 파일이 없으면(또는 --no-wiki) 빈 사전"""
    out = defaultdict(dict)
    path = WIKI / '인적사항.csv'
    if path.exists() and not NO_WIKI:
        with path.open(encoding='utf-8-sig', newline='') as f:
            for r in csv.DictReader(f):
                out[r['선수ID']][r['항목']] = (r['값'], r['출처'])
    return out


# 세부 포지션: IF는 내야 칸으로만, OF는 외야 칸으로만 좁힌다
WIKI_POS_NARROW = {'IF': {'1B', '2B', '3B', 'SS'}, 'OF': {'LF', 'CF', 'RF'}}
WIKI_POS_CODES = {'C', '1B', '2B', '3B', 'SS', 'LF', 'CF', 'RF', 'DH', 'IF', 'OF', 'P'}


def fill_players_from_wiki(players):
    """선수 마스터의 빈 생년·투타·언더핸드·세부 포지션을 위키 값으로 채운다. 채운 수를 돌려준다"""
    prof = wiki_profile()
    n = defaultdict(int)
    for p in players:
        w = prof.get(p['id'])
        if not w:
            continue
        if not p.get('birthYear') and '생년월일' in w:
            m = re.match(r'(\d{4})', w['생년월일'][0])
            if m:
                p['birthYear'] = int(m.group(1))
                n['생년'] += 1
        if not p.get('throws') and w.get('투', ('',))[0] in ('R', 'L'):
            p['throws'] = w['투'][0]
            n['투'] += 1
            if w.get('언더핸드', ('',))[0] == '예' and not p.get('underhand'):
                p['underhand'] = True
                n['언더핸드'] += 1
        if not p.get('bats') and w.get('타', ('',))[0] in ('R', 'L', 'S'):
            p['bats'] = w['타'][0]
            n['타'] += 1
        code = w.get('포지션 코드', ('',))[0]
        cur = p.get('pos')
        if code in WIKI_POS_CODES and code != cur and (cur is None or code in WIKI_POS_NARROW.get(cur, ())):
            # 투수·야수 구분과 맞을 때만 (투수에게 야수 칸, 야수에게 'P'를 붙이지 않는다)
            if cur is not None or (code == 'P') == (p['kind'] == 'P') or p['kind'] == 'BP':
                p['pos'] = code
                n['포지션'] += 1
    return n


# 위키 기록 항목 이름 → 시즌 행 필드 (tools/wiki/s5_merge.py의 FIELD_KO·PIT_KO와 같아야 한다)
WIKI_BAT_KO = {'경기': 'g', '타석': 'pa', '타수': 'ab', '득점': 'r', '안타': 'h', '2루타': 'd', '3루타': 't', '홈런': 'hr', '타점': 'rbi',
               '도루': 'sb', '도루실패': 'cs', '볼넷': 'bb', '몸에 맞는 공': 'hbp', '삼진': 'so', '병살타': 'gdp', '희생플라이': 'sf',
               '희생번트': 'sac'}
WIKI_PIT_KO = {'경기': 'g', '선발': 'gs', '승': 'w', '패': 'l', '세이브': 'sv', '홀드': 'hld', '아웃(이닝×3)': 'outs', '상대 타자': 'tbf',
               '자책점': 'er', '실점': 'r', '피안타': 'h', '피홈런': 'hr', '볼넷': 'bb', '몸에 맞는 공': 'hbp', '탈삼진': 'so'}
# 위키 표의 구단 이름 → 계보 (계보 열이 없는 옛 CSV용. tools/wiki/common.py의 FRANCHISE_KEYS와 같다. 위에서부터 처음 맞는 것)
WIKI_FRANCHISE_KEYS = [
    ('MBC', 'LG'), ('LG', 'LG'), ('삼성', '삼성'), ('롯데', '롯데'), ('해태', 'KIA'), ('KIA', 'KIA'), ('기아', 'KIA'),
    ('OB', '두산'), ('두산', '두산'), ('삼미', '현대계'), ('청보', '현대계'), ('태평양', '현대계'), ('현대', '현대계'),
    ('빙그레', '한화'), ('한화', '한화'), ('쌍방울', '쌍방울'), ('SSG', 'SSG'), ('SK', 'SSG'), ('우리', '히어로즈'),
    ('서울 히어로즈', '히어로즈'), ('히어로즈', '히어로즈'), ('넥센', '히어로즈'), ('키움', '히어로즈'), ('NC', 'NC'),
    ('KT', 'KT'), ('kt', 'KT'),
]


def wiki_franchise(text):
    for k, v in WIKI_FRANCHISE_KEYS:
        if re.search(r'(?<![A-Za-z])' + re.escape(k) + r'(?![A-Za-z])', text or ''):
            return v
    return None


def fill_records_from_wiki(bat, pit, teams):
    """기록_채움후보.csv로 있는 시즌 행의 빈 칸(None)만 채운다. 행이 없는 후보는 더하지 않는다.
    한 해 같은 계보에 그 선수 행이 둘 이상이면(시즌 중 같은 계보 안 이동) 어느 행인지 몰라 건너뛴다. (채움, 셈) 을 돌려준다"""
    path = WIKI / '기록_채움후보.csv'
    n = defaultdict(int)
    filled = defaultdict(int)  # (구분, 항목) → 칸
    if not path.exists() or NO_WIKI:
        return filled, n
    bi = {f: i for i, f in enumerate(BAT_FIELDS)}
    pi = {f: i for i, f in enumerate(PIT_FIELDS)}
    done = []  # (행, 열, 투수 여부, (구분, 항목))
    with path.open(encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(f):
            if r['우리 값'] == '(행 없음)':
                n['행 없음'] += 1
                continue
            y, pid, is_pit = int(r['연도']), r['선수ID'], r['구분'] == '투수'
            key = (WIKI_PIT_KO if is_pit else WIKI_BAT_KO).get(r['항목'])
            fr = r.get('계보') or wiki_franchise(r['구단'])  # 계보 열: s5_merge가 구단 칸을 낱말로 읽은 값 (한 구단 줄만)
            if not key or not fr or not r['채울 값']:
                n['읽지 못함'] += 1
                continue
            fr_of = {t['name']: t['franchise'] for t in teams.get(y, [])}
            rows_ = [row for row in (pit if is_pit else bat).get(y, []) if row[0] == pid and fr_of.get(row[1]) == fr]
            if len(rows_) != 1:
                n['행 여럿·못 찾음'] += 1
                continue
            col = (pi if is_pit else bi)[key]
            if rows_[0][col] is not None:
                n['이미 있음'] += 1  # 하드힛 등으로 이미 채워진 칸
                continue
            rows_[0][col] = int(float(r['채울 값']))
            done.append((rows_[0], col, is_pit, (r['구분'], r['항목'])))
    # 채운 뒤 타석과 맞지 않는 타자 행(타수+볼넷+사구 > 타석, validate_data.py와 같은 검사)은 그 행에 채운 칸을 모두 되돌린다
    bad = {id(row) for row, _, is_pit, _ in done
           if not is_pit and (row[bi['ab']] or 0) + (row[bi['bb']] or 0) + (row[bi['hbp']] or 0) > (row[bi['pa']] or 0)}
    for row, col, _, label in done:
        if id(row) in bad:
            row[col] = None
            n['타석과 맞지 않아 되돌림'] += 1
        else:
            filled[label] += 1
            n['채움'] += 1
    return filled, n


# 위키 병역 글 → 복무 형태. 위에서부터 처음 맞는 것 (임시 규칙, 2026-10-23)
WIKI_SERVICE = [
    ('exempt', r'예술체육요원|예술 ?체육 ?요원|병역 ?면제|면제|전시근로역|제2국민역'),
    ('상무', r'상무|국군체육부대'),
    ('경찰', r'경찰'),
    ('성무', r'성무'),  # 공군 야구단 (1970~80년대)
    ('reserve', r'공익|사회복무|(?<!수도)방위|보충역|산업기능요원'),
    ('active', r'현역|육군|해군|공군|해병|병장|전역|상근|사단|사령부|군단|여단|연대|대대|군필'),
]


def apply_record_fixes(bat, pit, teams):
    """기록_고침.csv (tools/wiki/resolve_209.py: 위키 209칸 중 신뢰 순서로 고칠 칸)를 시즌 행에 적용한다.
    저장소 값이 지금 행 값과 같을 때만 고친다 (원자료가 바뀌었으면 건너뛴다). (고친 칸 목록, 건너뛴 칸 목록)을 돌려준다"""
    path = WIKI / '기록_고침.csv'
    done, skip = [], []
    if not path.exists() or NO_WIKI:
        return done, skip
    bi = {f: i for i, f in enumerate(BAT_FIELDS)}
    pi = {f: i for i, f in enumerate(PIT_FIELDS)}
    with path.open(encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(f):
            y, pid, is_pit = int(r['연도']), r['선수ID'], r['구분'] == '투수'
            key = (WIKI_PIT_KO if is_pit else WIKI_BAT_KO).get(r['항목'])
            fr_of = {t['name']: t['franchise'] for t in teams.get(y, [])}
            rows_ = [row for row in (pit if is_pit else bat).get(y, []) if row[0] == pid and fr_of.get(row[1]) == r['계보']]
            col = (pi if is_pit else bi).get(key)
            label = f"{r['이름']}({pid}) {y} {r['구분']} {r['항목']} {r['저장소 값']} → {r['고칠 값']}"
            if col is None or len(rows_) != 1 or rows_[0][col] != int(r['저장소 값']):
                skip.append(label)
                continue
            rows_[0][col] = int(r['고칠 값'])
            done.append(label)
    return done, skip


def write_wiki_report(players_n, records, record_n, mil, special, fixes=((), ())):
    """reports/wiki-fill.md: 위키 값으로 채운 수 (build_data.py가 쓴다)"""
    L = ['# 위키 가공값 반영 결과', '',
         '`python tools/build_data.py`가 만든다. `data-src/wiki/`(위키데이터·위키백과·나무위키 가공값)로 우리 자료의 **빈 칸만** 채운다.',
         '우리 값과 다른 칸은 바꾸지 않는다 (`reports/wiki-diff.md`, `reports/wiki-diff-209.md`).', '',
         '## 선수 마스터', '', '| 항목 | 채운 선수 |', '|---|---:|']
    for k in ('생년', '투', '타', '언더핸드', '포지션'):
        L.append(f'| {k} | {players_n.get(k, 0)} |')
    L += ['', '## 병역 (military.json)', '',
          f"실제 자료(자료집 프로필 상무·경찰, 병역_면제명단.csv) {mil['real']}명. 그 밖의 국내 선수 중 위키 병역 칸이 있는 선수 {mil['wiki']}명을 `src: wiki`로 더했다. "
          f"읽지 못한 글 {len(mil['skip'])}개:", '']
    L += [f'- {x}' for x in mil['skip']]
    L += ['', '## 시즌 기록 빈 칸', '', '| 구분 | 항목 | 채운 칸 |', '|---|---|---:|']
    for (kind, item), v in sorted(records.items(), key=lambda kv: -kv[1]):
        L.append(f'| {kind} | {item} | {v} |')
    L += ['', '후보 처리: ' + ', '.join(f'{k} {v}' for k, v in sorted(record_n.items())) +
          ". '행 없음'은 우리 자료에 그 해 행이 없는 후보로, 선수 구성(저장 호환)이 바뀌어 더하지 않았다.", '',
          f"특별 엔트리(은퇴식 규칙) 연도별 인원: {sum(len(v) for v in special.values())}명 ({len(special)}개 연도).", '']
    L += ['## 위키 209칸 고침 (`data-src/wiki/기록_고침.csv`, `reports/wiki-209-resolve.md`)', '',
          f'하드힛·KBO 공식 자료가 둘 다 위키 값과 같은 칸만 고친다. 고침 {len(fixes[0])}칸, 건너뜀 {len(fixes[1])}칸.', '']
    L += [f'- 고침: {x}' for x in fixes[0]] + [f'- 건너뜀 (행이 없거나 저장소 값이 바뀜): {x}' for x in fixes[1]]
    L.append('')
    WIKI_REPORT.write_text('\n'.join(L) + '\n', encoding='utf-8', newline='\n')
    print(f"위키: 선수 {dict(players_n)}, 기록 {sum(records.values())}칸, 병역 {mil['wiki']}명 → {WIKI_REPORT}")


def wiki_military(table, players, seen):
    """실제 병역 자료가 없는 국내 선수에게 위키 병역 칸을 더한다 (table을 고친다). 더한 수와 읽지 못한 글 목록을 돌려준다.
    - 면제(예술체육요원·병역면제·전시근로역): exempt.year = 글 속 해, 없으면 첫 1군 시즌 (임시값: 언제 면제됐는지 모름)
    - 상무·경찰·성무: served.unit, 현역·보충역(공익·사회복무·방위·산업기능요원): served.unit '현역'·'보충역' + kind.
      end = 글 속 마지막 해 (없으면 모름). 미필·불명은 넣지 않는다"""
    prof = wiki_profile()
    added, skip = 0, []
    for p in players:
        if p.get('foreign') or p['id'] in table or '병역' not in prof.get(p['id'], {}):
            continue
        text = prof[p['id']]['병역'][0]
        if re.search(r'미필|불명|입대 예정', text):  # '복무 중'은 글 속 마지막 해를 전역 예정 해로 본다
            skip.append(f"{p['name']}({p['id']}): {text[:40]}")
            continue
        kind = next((k for k, pat in WIKI_SERVICE if re.search(pat, text)), None)
        if not kind:
            skip.append(f"{p['name']}({p['id']}): {text[:40]}")
            continue
        years = [int(y) for y in re.findall(r'(?<!\d)(19[4-9]\d|20[0-3]\d)(?!\d)', text)]
        if kind == 'exempt':
            year = max(years) if years else (p.get('first') or min(seen.get(p['id'], {}) or [None]) or None)
            if not year:
                skip.append(f"{p['name']}({p['id']}): {text[:40]}")
                continue
            event = '예술체육요원' if '요원' in text else '병역 면제'
            table[p['id']] = {'src': 'wiki', 'exempt': {'year': year, 'event': event}}
        else:
            unit = {'active': '현역', 'reserve': '보충역'}.get(kind, kind)
            # end는 실제 자료와 같은 뜻(복무로 빠진 마지막 시즌)으로 맞춘다: 전역한 해에 1군 기록이 있으면(시즌 중 전역, 방위병) 그 앞 해로.
            # 그대로 두면 그 해 시작 게임에서 '시작 뒤 복무'로 잡혀 다시 입대한다
            end = max(years) if years else None
            while end and end in seen.get(p['id'], {}) and end > min(years):
                end -= 1
            served = {'unit': unit, **({'end': end} if end else {}),
                      **({'kind': kind} if kind in ('active', 'reserve') else {})}
            table[p['id']] = {'src': 'wiki', 'served': served}
        added += 1
    return added, skip


def hardhit_csv(name):
    path = HARDHIT / name
    if not path.exists():
        return []
    with path.open(encoding='utf-8', newline='') as f:
        return list(csv.DictReader(f))


def fill_running(bat):
    """하드힛 도루·도루실패·병살타로 자료집의 빈 값을 채운다. 값이 이미 있는데 다르면 하드힛 값으로 바꾼다
    (2026-10-22 사용자 결정: 같은 값이 서로 다르면 하드힛 값을 쓴다).
    → 보고서 줄 목록"""
    hh = {(int(r['연도']), r['선수ID']): r for r in hardhit_csv('bat_running.csv')}
    if not hh:
        return []
    cols = (('sb', '도루'), ('cs', '도루실패'), ('gdp', '병살타'))
    idx = {f: BAT_FIELDS.index(f) for f, _ in cols}
    filled, differ, same = defaultdict(int), defaultdict(int), defaultdict(int)
    filled_y = defaultdict(lambda: defaultdict(int))
    examples = defaultdict(list)
    for y, rows_ in bat.items():
        for row in rows_:
            h = hh.get((y, row[0]))
            if not h:
                continue
            for f, k in cols:
                if h[k] == '':
                    continue
                v = int(h[k])
                cur = row[idx[f]]
                if cur is None:
                    row[idx[f]] = v
                    filled[f] += 1
                    filled_y[y][f] += 1
                elif cur != v:
                    row[idx[f]] = v
                    differ[f] += 1
                    if len(examples[f]) < 20:
                        examples[f].append(f'{y} {row[0]} 자료집 {cur} / 하드힛 {v}')
                else:
                    same[f] += 1
    out = ['## 타자 도루·도루실패·병살타 (bat_running.csv)', '',
           '빈 값은 하드힛 값으로 채우고, 값이 있는데 다르면 하드힛 값으로 바꾼다 (2026-10-22 사용자 결정).', '',
           '| 항목 | 채움 | 같음 | 하드힛 값으로 바꿈 |', '|---|---:|---:|---:|']
    out += [f'| {k} | {filled[f]} | {same[f]} | {differ[f]} |' for f, k in cols]
    out += ['', '연도별 채움 (도루 / 도루실패 / 병살타):', '']
    out += [f'- {y}: {d["sb"]} / {d["cs"]} / {d["gdp"]}' for y, d in sorted(filled_y.items())]
    for f, k in cols:
        if examples[f]:
            out += ['', f'{k} 하드힛 값으로 바꾼 것 (연도 선수ID):', ''] + [f'- {e}' for e in examples[f]]
    print(f'하드힛 채움: 도루 {filled["sb"]}, 도루실패 {filled["cs"]}, 병살타 {filled["gdp"]} / 하드힛 값으로 바꿈 {dict(differ)}')
    return out


FIELD_POS = ('C', '1B', '2B', '3B', 'SS', 'LF', 'CF', 'RF')


def pos_2009(row_pos, hh_pos):
    """2009년 타자 주 포지션 (하드힛 2009년 파일 없음). 2026-10-22 사용자 결정 (나):
    자료집 시즌 행에 세부 포지션이 있으면 그 값, 없으면(빈 값·내야수·외야수) 하드힛 2008·2010 중 있는 해
    (둘 다 있고 다르면 2008), 둘 다 없으면 하드힛에 나온 가장 가까운 해(같은 거리면 앞 해).
    hh_pos: 그 선수의 {연도: 하드힛 주 포지션}. → (포지션, 출처) 출처는 '자료집' / '빌림' / None(정할 수 없음)"""
    if row_pos in FIELD_POS or row_pos == 'DH':
        return row_pos, '자료집'
    near = [y for y in (2008, 2010) if y in hh_pos]
    if near:
        return hh_pos[near[0]], '빌림'
    if hh_pos:
        y = min(hh_pos, key=lambda k: (abs(k - 2009), k))
        return hh_pos[y], '빌림'
    return None, None


def fill_positions(bat, kind_of):
    """하드힛 그 시즌 주 포지션(position.csv, 2003~)으로 타자 시즌 행의 포지션을 채우고 바로잡는다 (2026-10-22 사용자 결정).
    - 하드힛에 나온 선수-시즌: 하드힛 값 (자료집 값이 다르면 하드힛 값으로 바꾼다). 한 해 여러 구단 행은 모두 같은 값.
    - 하드힛 해인데 어느 포지션 파일에도 없는 타자 [B]: 직전 시즌 하드힛 주 포지션 → 그 앞 가장 가까운 시즌
      → (앞 시즌이 없으면) 뒤의 가장 가까운 시즌 → 하드힛 전체에 없으면 DH. 타석 0인 행은 DH 대신 자료집 값을 둔다.
    - 2009년(하드힛 파일 없음): pos_2009.
    posSrc 열: 'H' 하드힛 확인값, 'E' 추정값(2009, [B]), None 자료집 그대로(2002년 이전 등). 엔진은 'H'만 규정타석 판정에 쓴다.
    → (보고서 줄 목록, 빈 포지션 선수 목록용 정보)"""
    hh = {}
    by_player = defaultdict(dict)
    for r in hardhit_csv('position.csv'):
        y, pid = int(r['연도']), r['선수ID']
        hh[(y, pid)] = r['주 포지션']
        by_player[pid][y] = r['주 포지션']
    if not hh:
        return [], []
    hh_years = sorted({y for y, _ in hh})
    span = range(hh_years[0], hh_years[-1] + 1)
    pi, si, ai = BAT_FIELDS.index('pos'), BAT_FIELDS.index('posSrc'), BAT_FIELDS.index('pa')
    stat = defaultdict(lambda: defaultdict(int))
    examples = []
    unresolved = []  # (사유, 연도, 선수ID) — 하드힛 해인데 하드힛 포지션을 끝내 얻지 못한 선수-시즌
    for y in span:
        pa_of = defaultdict(int)
        for row in bat.get(y, []):
            pa_of[row[0]] += row[ai] or 0
        for row in bat.get(y, []):
            pid, cur = row[0], row[pi]
            if kind_of.get(pid) == 'P':
                continue
            S = stat[y]
            if y not in hh_years:
                v, how = pos_2009(cur, by_player.get(pid, {}))
                if how == '자료집':
                    S['2009 자료집'] += 1
                elif how == '빌림':
                    S['2009 빌림'] += 1
                    row[pi], row[si] = v, 'E'
                else:
                    S['2009 못 정함'] += 1
                    unresolved.append(('2009 자료 없음, 다른 시즌 확인 불가', y, pid))
                if how == '자료집':
                    row[si] = 'E'
                continue
            v = hh.get((y, pid))
            if v:
                if cur is None or cur in ('IF', 'OF'):
                    S['채움'] += 1
                elif cur != v:
                    S['바꿈'] += 1
                    if len(examples) < 40:
                        examples.append(f'{y} {pid} 자료집 {cur} → 하드힛 {v}')
                else:
                    S['같음'] += 1
                row[pi], row[si] = v, 'H'
                continue
            # [B] 어느 포지션 파일에도 없는 타자
            seen = by_player.get(pid, {})
            before = [k for k in seen if k < y]
            after = [k for k in seen if k > y]
            if before:
                v = seen[max(before)]
            elif after:
                v = seen[min(after)]
            elif pa_of[pid] > 0:
                v = 'DH'
                unresolved.append(('하드힛 전 시즌 포지션 없음', y, pid))
            else:
                S['[B] 타석 0 그대로'] += 1
                continue
            S['[B] DH' if v == 'DH' else '[B] 다른 해에서'] += 1
            row[pi], row[si] = v, 'E'
    keys = ['채움', '바꿈', '같음', '[B] 다른 해에서', '[B] DH', '[B] 타석 0 그대로', '2009 자료집', '2009 빌림', '2009 못 정함']
    tot = {k: sum(stat[y][k] for y in stat) for k in keys}
    print('하드힛 주 포지션: ' + ', '.join(f'{k} {tot[k]}' for k in keys if tot[k]))
    out = ['', '## 타자 주 포지션 (position.csv)', '',
           '하드힛 그 시즌 주 포지션으로 채우고, 자료집 값이 다르면 하드힛 값으로 바꾼다 (2026-10-22 사용자 결정). 단위: 타자 시즌 행 (한 해 여러 구단이면 행마다).',
           '- 채움: 자료집이 빈 값이거나 내야수·외야수뿐이던 행. 바꿈: 세부 포지션이 하드힛과 달라 바꾼 행.',
           '- [B]: 하드힛 해인데 어느 포지션 파일에도 없는 타자 (지명타자 후보). 직전 시즌 → 앞의 가장 가까운 시즌 → 뒤의 가장 가까운 시즌의 하드힛 주 포지션, 하드힛 전체에 없으면 DH.',
           '- 2009년 (하드힛 파일 없음, 사용자 결정 (나)): 자료집 세부 포지션이 있으면 그 값, 없으면 2008·2010 하드힛(다르면 2008), 그것도 없으면 가장 가까운 해. 2009년 값은 규정타석 판정에 쓰지 않는다.', '',
           '| 연도 | ' + ' | '.join(keys) + ' |', '|---|' + '---:|' * len(keys)]
    out += [f'| {y} | ' + ' | '.join(str(stat[y][k]) for k in keys) + ' |' for y in sorted(stat)]
    out += [f'| 합계 | ' + ' | '.join(str(tot[k]) for k in keys) + ' |']
    if examples:
        out += ['', '바꾼 예 (앞 40건):', ''] + [f'- {e}' for e in examples]
    return out, unresolved


NO_POSITION_XLSX = ROOT / 'reports' / 'hardhit-no-position.xlsx'
KO_POS = {'C': '포수', '1B': '1루수', '2B': '2루수', '3B': '3루수', 'SS': '유격수', 'LF': '좌익수', 'CF': '중견수', 'RF': '우익수',
          'DH': '지명타자', 'IF': '내야수', 'OF': '외야수'}


def write_no_position(unresolved, bat, players):
    """하드힛 해(2003~)에 타석이 있는데 하드힛 주 포지션을 끝내 얻지 못한 타자 목록 → reports/hardhit-no-position.xlsx.
    한 선수가 두 사유에 다 걸리면 사유마다 한 줄. 타석 합 많은 순"""
    if not unresolved:
        return
    from datetime import date
    from openpyxl import Workbook
    from openpyxl.styles import Font
    pmap = {p['id']: p for p in players}
    pi, si, ai, ti = (BAT_FIELDS.index(k) for k in ('pos', 'posSrc', 'pa', 'team'))
    reasons = defaultdict(set)
    for why, _, pid in unresolved:
        reasons[pid].add(why)
    out = []
    for pid, whys in reasons.items():
        seasons = defaultdict(lambda: [0, [], defaultdict(int)])  # 연도 → 타석, 팀, 포지션별 타석
        for y in range(2003, max(bat) + 1):
            for row in bat.get(y, []):
                if row[0] == pid and (row[ai] or 0) > 0:
                    s = seasons[y]
                    s[0] += row[ai]
                    s[1].append(row[ti])
                    s[2][row[pi]] += row[ai]
        pos_pa = defaultdict(int)
        for s in seasons.values():
            for q, v in s[2].items():
                pos_pa[q] += v
        game_pos = max(pos_pa.items(), key=lambda kv: kv[1])[0] if pos_pa else None
        m = pmap.get(pid, {})
        for why in sorted(whys):
            out.append([pid, m.get('name', ''), '실존' if m.get('real', True) else '가상',
                        ', '.join(str(y) for y in sorted(seasons)),
                        ', '.join(f'{y} {"/".join(dict.fromkeys(seasons[y][1]))}' for y in sorted(seasons)),
                        sum(s[0] for s in seasons.values()),
                        KO_POS.get(m.get('pos'), m.get('pos') or ''), KO_POS.get(game_pos, game_pos or ''), why])
    out.sort(key=lambda r: (-r[5], r[0]))
    wb = Workbook()
    ws = wb.active
    ws.title = '포지션 없음'
    ws.append([f'만든 날짜 {date.today().isoformat()} · python tools/build_data.py가 다시 만든다'])
    ws.append(['기준: 하드힛 포지션 자료가 있는 해(2003~2026, 2009년 파일 없음)에 타석이 있는데 하드힛 주 포지션을 얻지 못한 타자. '
               '"하드힛 전 시즌 포지션 없음" = 어느 해 포지션 파일에도 없어 지명타자로 둠, '
               '"2009 자료 없음, 다른 시즌 확인 불가" = 2009년 자료집에 세부 포지션이 없고 다른 해 하드힛에도 없음. '
               '게임 주 포지션은 하드힛 해 시즌 행에서 타석이 가장 많은 포지션 (시작 연도마다 엔진이 직전 시즌으로 다시 고른다).'])
    ws['A1'].font = Font(bold=True)
    ws.append([])
    head = ['선수ID', '이름', '실존 여부', '하드힛 타석 있는 해', '그 해들의 팀', '타석 합', '선수마스터 포지션', '게임 주 포지션', '사유']
    ws.append(head)
    for c in ws[4]:
        c.font = Font(bold=True)
    for r in out:
        ws.append(r)
    for col, w in zip('ABCDEFGHI', (9, 10, 8, 22, 34, 8, 14, 12, 34)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A5'
    NO_POSITION_XLSX.parent.mkdir(exist_ok=True)
    wb.save(NO_POSITION_XLSX)
    cnt = defaultdict(int)
    for r in out:
        cnt[r[8]] += 1
    print(f'포지션 없음 목록: {len(out)}줄 ({", ".join(f"{k} {v}" for k, v in cnt.items())}) → {NO_POSITION_XLSX.name}')


def fill_gmli(pit):
    """하드힛 투수 situation의 gmLI(등판 때 평균 레버리지, 2010~)를 시즌 투수 행의 gmli에 넣는다. 자료집에는 없는 값이다.
    시즌 중 이적한 투수는 하드힛 값이 시즌 합계 하나라 그 해 모든 행에 같은 값을 넣는다. → 보고서 줄 목록"""
    gi = PIT_FIELDS.index('gmli')
    hh = {}
    for r in hardhit_csv('clutch_pit.csv'):
        if r['gmLI'] != '' and float(r['gmLI']) > 0:
            hh[(int(r['연도']), r['선수ID'])] = round(float(r['gmLI']), 2)
    if not hh:
        return []
    filled = defaultdict(int)
    for y, rows_ in pit.items():
        for row in rows_:
            v = hh.get((y, row[0]))
            if v is not None:
                row[gi] = v
                filled[y] += 1
    print(f'하드힛 gmLI: {sum(filled.values())}행')
    return ['## 투수 gmLI (clutch_pit.csv)', '', '등판 때 평균 레버리지. 자료집에 없는 값이라 넣기만 한다 (0인 행은 값 없음으로 본다).', '',
            ', '.join(f'{y}: {c}' for y, c in sorted(filled.items())), '']


# 선발 등판 추정 (src/engine/params.ts estimatedStarts와 같은 식): 경기당 아웃이 구원 3.5 ~ 선발 14 사이 어디인가
GS_EST_RELIEF_OUTS, GS_EST_START_OUTS = 3.5, 14


def fill_pitch_roles(pit):
    """자료집에 보직(role)·선발 등판(gs)이 빈 투수 행을 하드힛 투수 classic(pit_classic.csv)으로 채운다 (2026년은 296명이 비어 있다).
    하드힛 classic에는 선발 등판 열이 없어 gs는 경기당 아웃 추정에 퀄리티 스타트를 하한으로 둔 값이다 (2015~2025 자료집 gs와 견준 오차 투수당 약 0.5번).
    하드힛 보직 라벨은 "선발"만 쓴다 (마무리·셋업 라벨은 해마다 1~16명뿐이라 안 쓴다, 기획안 1.3).
    role: 라벨 선발이거나 추정 선발 비율 0.5 이상이면 SP, 아니고 세이브 10개 이상이면 CL, 나머지 RP. 값이 있는 행은 건드리지 않는다.
    선발 등판 자료가 통째로 빠진 해만 채운다 (추정값이 실제 기록 사이에 섞이지 않게). → 보고서 줄 목록"""
    hh = {(int(r['연도']), r['선수ID']): r for r in hardhit_csv('pit_classic.csv')}
    if not hh:
        return []
    ix = {f: i for i, f in enumerate(PIT_FIELDS)}
    filled = defaultdict(lambda: [0, 0])
    for y, rows_ in pit.items():
        # 선발 등판 자료가 통째로 빠진 해(자료집 보조 열이 없는 2026 스냅샷)만 채운다. 다른 해의 빈 칸은 엔진이 경기 때 같은 식으로 추정한다
        if sum(r[ix['gs']] is None for r in rows_) < 0.9 * len(rows_):
            continue
        for row in rows_:
            h = hh.get((y, row[0]))
            if not h or (row[ix['role']] is not None and row[ix['gs']] is not None):
                continue
            g, outs = row[ix['g']] or 0, row[ix['outs']] or 0
            share = 0 if g <= 0 else min(1, max(0, (outs / g - GS_EST_RELIEF_OUTS) / (GS_EST_START_OUTS - GS_EST_RELIEF_OUTS)))
            qs = int(h['퀄스'] or 0)
            gs = min(g, max(qs, round(g * share)))
            if row[ix['gs']] is None:
                row[ix['gs']] = gs
                filled[y][1] += 1
            if row[ix['role']] is None:
                row[ix['role']] = 'SP' if h['보직'] == '선발' or share >= 0.5 else 'CL' if (row[ix['sv']] or 0) >= 10 else 'RP'
                filled[y][0] += 1
    print(f'하드힛 보직·선발 채움: {sum(v[0] for v in filled.values())}행 / {sum(v[1] for v in filled.values())}행')
    return ['## 투수 보직·선발 등판 (pit_classic.csv)', '',
            '선발 등판 자료가 통째로 빠진 해(2026 스냅샷)의 빈 보직·선발 등판만 채운다. 선발 등판은 **추정값**, 경기당 아웃 추정에 퀄리티 스타트를 하한으로 둔 값 (하드힛에 선발 등판 열이 없다).', '',
            '연도별 채움 (보직 / 선발 등판):', ''] + [f'- {y}: {a} / {b}' for y, (a, b) in sorted(filled.items())] + ['']


def fill_hands(players):
    """하드힛 투구손(좌/우/좌언/우언)으로 선수 마스터의 빈 throws·underhand를 채운다. 이미 있는데 다르면 하드힛 값으로 바꾼다.
    → 보고서 줄 목록"""
    seen = defaultdict(lambda: defaultdict(int))
    for r in hardhit_csv('pitch_hand.csv'):
        seen[r['선수ID']][(r['손'], r['언더'] == '1')] += 1
    if not seen:
        return []
    pmap = {p['id']: p for p in players}
    filled, same, varies = 0, 0, 0
    differ = []
    for pid, cnt in seen.items():
        p = pmap.get(pid)
        if not p:
            continue
        varies += len(cnt) > 1
        th, under = max(cnt.items(), key=lambda kv: kv[1])[0]
        if p.get('throws') is None:
            p['throws'] = th
            if under:
                p['underhand'] = True
            filled += 1
        elif p['throws'] != th or bool(p.get('underhand')) != under:
            differ.append(f'{p["name"]}({pid}) 자료집 {p["throws"]}{" 언더" if p.get("underhand") else ""} / 하드힛 {th}{" 언더" if under else ""}')
            p['throws'] = th
            if under:
                p['underhand'] = True
            else:
                p.pop('underhand', None)
        else:
            same += 1
    print(f'하드힛 투구손: 채움 {filled}, 같음 {same}, 하드힛 값으로 바꿈 {len(differ)}, 해마다 갈림 {varies}')
    return ['', '## 투수 투구손 (pitch_hand.csv)', '',
            f'하드힛 투수 {len(seen)}명: 빈 투타 채움 {filled}명, 같음 {same}명, 하드힛 값으로 바꿈 {len(differ)}명, 해마다 손 구분이 갈린 선수 {varies}명.', ''] + [f'- {d}' for d in differ]


def park_factors(teams):
    """하드힛 파크팩터(선수 행의 타석 가중 평균, 홈 경기 절반만 반영한 시즌 단위 값) → 연도 → 계보 → 홈 구장 득점 계수.

    홈 구장 계수 = (2 × 파크팩터 − 100) / 100 (원정 경기는 리그 평균 구장으로 본다). 해마다 구단 평균이 1이 되게 나눈다.
    자료가 빠진 해는 앞뒤 해의 평균, 한쪽만 있으면 그 값 (임시값). 지금은 빠진 해가 없다
    (2022는 원본 파일이 2020 자료라 한때 앞뒤 해 평균을 썼고, 2026-10-22 다시 받은 파일로 실제 값이 되었다).
    2015년 이전은 자료가 없어 넣지 않는다 (엔진에서 1 = 중립)."""
    raw = {(int(r['연도']), r['계보']): float(r['파크팩터']) for r in hardhit_csv('park.csv')}
    if not raw:
        return {}, []
    first = min(y for y, _ in raw)
    out, filled = {}, []
    for y in sorted(teams):
        if y < first:
            continue
        year = {}
        for t in teams[y]:
            fr = t['franchise']
            if (y, fr) in raw:
                pf = raw[(y, fr)]
            else:
                near = [raw[(y + d, fr)] for d in (-1, 1) if (y + d, fr) in raw]
                if not near:
                    continue
                pf = sum(near) / len(near)
                filled.append(f'{y} {fr}')
            year[fr] = (2 * pf - 100) / 100
        mean = sum(year.values()) / len(year)
        out[str(y)] = {fr: round(v / mean, 3) for fr, v in year.items()}
    print(f'하드힛 파크팩터: {len(out)}개 연도, 앞뒤 해로 채움 {len(filled)}건')
    lines = ['', '## 홈 구장 계수 (park.csv → meta.json parkFactors)', '',
             '홈 구장 계수 = (2 × 파크팩터 − 100) / 100, 해마다 구단 평균 1로 나눔. 앞뒤 해 평균으로 채운 것(임시값): ' + (', '.join(filled) or '없음'), '',
             '| 연도 | ' + ' | '.join(sorted(next(iter(out.values())))) + ' |', '|---|' + '---:|' * len(next(iter(out.values())))]
    for y, v in out.items():
        lines.append(f'| {y} | ' + ' | '.join(f'{v.get(fr, 1):.3f}' for fr in sorted(next(iter(out.values())))) + ' |')
    return out, lines


# 구장 표의 구장 → (계보들, 그 구장이 홈이 된 해). 개장 해는 src/league/ballparks.ts와 같다.
# 포항(삼성)·울산(롯데)은 제2구장이라 엔진에 없어 쓰지 않는다 (몇 경기씩만 열린다)
PARK_STADIUMS = {'대구': (['삼성'], 2016), '인천': (['SSG'], 2002), '부산': (['롯데'], 1986), '대전': (['한화'], 2025),
                 '수원': (['KT'], 2014), '창원': (['NC'], 2019), '광주': (['KIA'], 2014), '잠실': (['LG', '두산'], 1985),
                 '고척': (['히어로즈'], 2016)}
# 구성 지수를 100 쪽으로 당기는 정도: 로그 지수 × 경기 / (경기 + PARK_SHRINK_GAMES). 288경기(4시즌) 구장 0.67, 대전 73경기 0.34 (임시값)
PARK_SHRINK_GAMES = 144
# 좌타·우타 표를 전체 표 쪽으로 당기는 정도: (좌우 log 지수 − 전체 log 지수) × 반 경기 / (반 경기 + PARK_SIDE_SHRINK_GAMES) (임시값.
# 좌우 표는 표본이 전체의 절반쯤이다. 288경기 구장 0.5, 대전 0.2)
PARK_SIDE_SHRINK_GAMES = 144
# 사건별 득점 가치: 그 사건이 1개 더 날 때 엔진에서 느는 득점 (tools/park_run_value.ts, 2025년 6시드). 구성 배율이 내는 득점 계수를 셀 때만 쓴다.
# 흔한 선형 가중치(단타 0.74, 2루타 1.05, 홈런 1.68)를 쓰면 엔진 득점 비가 f보다 1.25배 크게 움직였다 (reports/park-check.md).
# 3루타는 잰 값(0.47)이 표본이 작아 2루타 값을 쓴다 (임시값)
PARK_RUN_VALUE = {'s1': 0.61, 'd2': 0.69, 't3': 0.69, 'hr': 1.67}
# 수준 맞춤에 쓰는 기본 나눔 (엔진 params.park의 hitExp·hrExp와 같다)
PARK_DEFAULT_EXP = {'s1': 0.4, 'd2': 0.4, 't3': 0.4, 'hr': 1.2}


def park_events(parks, league):
    """구장 표의 구성 모양 + 연도별 홈 구장 계수 f(수준) → 연도 → 계보 → 사건별 구장 배율 {s1, d2, t3, hr}.

    배율 = exp(w × log 구성 지수 + 기본 지수 × d). w는 경기 수로 당긴 정도, d는 그 해 리그 비율·득점 가치로 센 득점 계수가
    그 해 f와 같아지게 정한다 (모양은 표에서, 수준은 해마다 f에서). 볼넷·삼진 지수는 쓰지 않는다:
    표의 득점 지수와 견주면 안타·홈런 네 가지로 센 득점 계수가 더 잘 맞는다 (상관 0.96, 볼넷·삼진까지 넣으면 0.84).
    그 구장이 홈이 된 해부터만 쓴다. 표가 없는 구장(예전 구장)은 엔진이 기본 지수(f^hitExp, f^hrExp)를 쓴다."""
    rows = hardhit_csv('park_components.csv')
    if not rows or not parks:
        return {}, []
    col = {'s1': '단타', 'd2': '2루타', 't3': '3루타', 'hr': '홈런'}
    shape = {}  # 계보 → (개장 해, {사건: 당긴 log 지수}, 구장, 경기, {'L'|'R': {사건: 당긴 log 지수}})
    side_rows = {}
    for r in rows:
        if r['구장'] not in PARK_STADIUMS:
            continue
        if r['타자'] != 'ALL':
            side_rows[(r['구장'], r['타자'])] = r
            continue
        frs, since = PARK_STADIUMS[r['구장']]
        g = int(r['경기'])
        w = g / (g + PARK_SHRINK_GAMES)
        for fr in frs:
            shape[fr] = (since, {e: w * math.log(float(r[c]) / 100) for e, c in col.items()}, r['구장'], g, {})
    for fr, (since, base, st, g, sides) in shape.items():
        all_row = next(r for r in rows if r['구장'] == st and r['타자'] == 'ALL')
        for sd in ('L', 'R'):
            r = side_rows.get((st, sd))
            if not r:
                continue
            half = int(r['경기']) / 2
            ws = half / (half + PARK_SIDE_SHRINK_GAMES)
            sides[sd] = {e: base[e] + ws * (math.log(float(r[c]) / 100) - math.log(float(all_row[c]) / 100)) for e, c in col.items()}
    out = {}
    for y, year in parks.items():
        R = league[int(y)]['rates']
        rpa = league[int(y)]['r'] / league[int(y)]['pa']
        ev = {}
        for fr, f in year.items():
            if fr not in shape or int(y) < shape[fr][0]:
                continue
            base, sides = shape[fr][1], shape[fr][4]

            def mult(d, b=base):
                return {e: math.exp(b[e] + PARK_DEFAULT_EXP[e] * d) for e in b}

            def run(d):
                m = mult(d)
                return 1 + sum(PARK_RUN_VALUE[e] * R[e] * (m[e] - 1) for e in m) / rpa

            lo, hi = -1.0, 1.0
            for _ in range(60):
                mid = (lo + hi) / 2
                lo, hi = (mid, hi) if run(mid) < f else (lo, mid)
            d = (lo + hi) / 2
            ev[fr] = {e: round(v, 3) for e, v in mult(d).items()}
            # 좌타·우타: 모양만 좌우 표에서, 수준 맞춤 d는 전체와 같다
            for sd, b in sides.items():
                ev[fr][sd] = {e: round(v, 3) for e, v in mult(d, b).items()}
        if ev:
            out[y] = ev
    lines = ['', '## 사건별 구장 배율 (park_components.csv → meta.json parkEvents)', '',
             '구장 표(2022~2025 묶음)는 100 = 리그 평균인 **구장 자체 지수**다. 표의 득점 지수와 연도별 홈 구장 계수 f(2022~2025 평균)의 기울기 0.72·상관 0.96,',
             '하드힛 선수 행 파크팩터(홈 경기 절반을 섞은 시즌 값)와는 기울기 0.35라 절반 값이 아니다.',
             f'모양: 단타·2루타·3루타·홈런 지수를 log로, 경기 / (경기 + {PARK_SHRINK_GAMES})만큼 100 쪽으로 당김 (임시값). 수준: 그 해 f에 맞춤. 볼넷·삼진은 쓰지 않음.',
             f'좌타(L)·우타(R): 좌우 표의 전체 대비 log 차이를 반 경기 / (반 경기 + {PARK_SIDE_SHRINK_GAMES})만큼만 더한다 (임시값). 수준 맞춤은 전체와 같다.',
             '엔진: 좌타는 L, 우타는 R, 양타는 상대 투수 손의 반대쪽, 치는 손을 모르면 전체. 방향별 홈런(좌·중·우)은 엔진에 타구 방향이 없어 쓰지 않는다.',
             '제2구장 포항(삼성 14경기)·울산(롯데 20경기)은 엔진에 없어 쓰지 않는다.', '',
             '| 계보 | 구장 | 표 경기 | 쓰는 해 | 2025 f | 타자 | 2025 단타 | 2루타 | 3루타 | 홈런 |', '|---|---|---:|---|---:|---|---:|---:|---:|---:|']
    for fr, (since, _, st, g, _s) in sorted(shape.items()):
        yrs = [y for y in out if fr in out[y]]
        e = out.get('2025', {}).get(fr) or {}
        for label, x in (('전체', e), ('좌타', e.get('L') or {}), ('우타', e.get('R') or {})):
            head = f'| {fr} | {st} | {g} | {yrs[0] if yrs else "-"}~{yrs[-1] if yrs else ""} | {parks.get("2025", {}).get(fr, 1):.3f} |' if label == '전체' else '| | | | | |'
            lines.append(f'{head} {label} | ' + ' | '.join(f'{x.get(k, 1):.3f}' for k in ('s1', 'd2', 't3', 'hr')) + ' |')
    print(f'사건별 구장 배율: {len(out)}개 연도, {sum(len(v) for v in out.values())}건')
    return out, lines


def profiles_of(wb):
    return {str(r['선수ID']): r['출신교·경력(2026 프로필)'] for r in rows(wb['선수마스터']) if r['출신교·경력(2026 프로필)']}


def main():
    if '--military-only' in sys.argv:
        # 병역 자료만 다시 만든다 (다른 JSON은 그대로라 저장 호환에 영향 없음)
        build_military(profiles_of(load_workbook(SRC, read_only=True, data_only=True)), OUT)
        return
    if '--drafts-only' in sys.argv:
        # 지명 자료만 다시 만든다 (선수 마스터는 players.json에서 읽는다)
        dump_drafts(draft_rows(), json.loads((OUT / 'players.json').read_text(encoding='utf-8'))['players'])
        return
    if '--managers-only' in sys.argv:
        # 감독 CSV만 다시 반영한다 (시즌 JSON에서 운영 경향을 계산한다)
        build_managers(OUT)
        return
    if '--ppa-only' in sys.argv:
        # 선수별 P/PA만 다시 반영한다 (가공 CSV에서 읽는다, 저장 호환은 그대로)
        print('선수별 P/PA 타자 %d·투수 %d 시즌' % build_ppa())
        return
    if '--running-only' in sys.argv:
        # 선수별 추가 진루·번트만 다시 반영한다 (가공 CSV에서 읽는다, 저장 호환은 그대로)
        print('선수별 추가 진루·번트 %d명' % build_running())
        return
    if '--plate-only' in sys.argv:
        # 선수별 볼카운트 사슬 실측만 다시 반영한다 (가공 CSV에서 읽는다, 저장 호환은 그대로)
        print('선수별 사슬 실측 타자 %d·투수 %d 시즌' % build_plate())
        return
    if '--homes-only' in sys.argv:
        # 연고지 CSV만 다시 반영한다 (다른 JSON과 저장 호환은 그대로)
        print(f'연고지 {len(build_homes())}행')
        return
    if '--rules-only' in sys.argv:
        # 연도별 규칙(meta.json rules)과 규칙 변경 CSV(rule_changes.json)만 다시 만든다 (엑셀을 읽지 않는다. 선수 월드 지문과 무관해 저장 호환은 그대로)
        meta = json.loads((OUT / 'meta.json').read_text(encoding='utf-8'))
        meta['rules'] = {y: rules_for(y) for y in meta['years']}
        (OUT / 'meta.json').write_text(json.dumps(meta, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
        print(f'연도별 규칙 {len(meta["rules"])}개 연도')
        build_rule_changes(OUT)
        return
    if '--schedule-only' in sys.argv:
        # 일정 규칙 CSV만 다시 반영한다 (다른 JSON과 저장 호환은 그대로)
        print(f'일정 규칙 {len(build_schedule()[0])}개 연도 + 자료 없는 해 규칙')
        return
    if '--history-only' in sys.argv:
        # 수상·구장·올스타전 CSV만 다시 반영한다 (다른 JSON과 저장 호환은 그대로)
        build_history(OUT)
        return
    if '--defense-only' in sys.argv:
        # 수비 이닝·평판 CSV만 다시 반영한다 (시즌 JSON에서 타석을 읽는다. 다른 JSON과 저장 호환은 그대로)
        build_defense(OUT)
        return
    if '--pitches-only' in sys.argv:
        # 구종 CSV(visualbaseball·위키)만 다시 반영한다 (다른 JSON과 저장 호환은 그대로)
        build_pitches(OUT)
        return
    if '--traits-only' in sys.argv:
        # 특수능력 CSV만 다시 반영한다 (엑셀을 읽지 않으므로 다른 JSON과 저장 호환은 그대로)
        ids = {p['id'] for p in json.loads((OUT / 'players.json').read_text(encoding='utf-8'))['players']}
        dump_traits(build_traits(ids))
        return
    print(f'읽는 중: {SRC}')
    wb = load_workbook(SRC, read_only=True, data_only=True)

    # ---- 선수 마스터
    players = []
    kind_of = {}
    for r in rows(wb['선수마스터']):
        th, ba, under = hands(r['투타'])
        kind = {'타자': 'B', '투수': 'P'}.get(r['구분'], 'BP')
        pid = str(r['선수ID'])
        kind_of[pid] = kind
        foreign = {'Y': True, 'N': False}.get(r['외국인'])
        players.append({
            'id': pid, 'name': r['선수명'], 'kind': kind, 'pos': POS.get(r['주포지션']),
            'throws': th, 'bats': ba, 'underhand': under or None,
            'birthYear': r['출생연도'], 'foreign': foreign,
            'entryYear': r['입단연도(추정)'], 'entryExact': r['입단연도(보조자료)'],
            'first': r['첫 1군 시즌'], 'last': r['마지막 1군 시즌'],
            'school': school_of(r['출신교·경력(2026 프로필)']),
            'real': True,
        })
    for p in players:
        for k in [k for k, v in p.items() if v is None]:
            del p[k]

    wiki_players = fill_players_from_wiki(players)
    hand_report = fill_hands(players)

    # ---- 외국인 개막 명단 (CSV)
    foreign_opening = defaultdict(lambda: defaultdict(list))
    asia_ids = set()
    if FOREIGN_OPENING.exists():
        with FOREIGN_OPENING.open(encoding='utf-8-sig', newline='') as f:
            for r in csv.DictReader(f):
                # 계약구분 열이 없던 예전 파일은 구분만 본다
                asia = r['구분'] == '아시아쿼터' or r.get('계약구분') == '아시아쿼터'
                if asia:
                    asia_ids.add(str(r['선수ID']))
                if r['구분'] == '제외':
                    continue
                entry = {'id': str(r['선수ID']), 'asia': asia, 'confirmed': r['확인'] in ('확인', '출처 확인')}
                if r.get('기록소속'):
                    entry['recordTeam'] = r['기록소속']
                # 실제 계약 금액 (달러, 있는 선수만): 연봉·계약금·출처
                if (r.get('연봉(달러)') or '').strip():
                    entry['salary'] = int(r['연봉(달러)'])
                    entry['bonus'] = int(r.get('계약금(달러)') or 0)
                    entry['source'] = (r.get('금액 출처') or '').strip()
                foreign_opening[str(r['연도'])][r['구단']].append(entry)
    for p in players:
        if p['id'] in asia_ids:
            p['asia'] = True

    # ---- 신인 지명 (baseballchart.kr 원본)
    drafts = draft_rows()

    # ---- 계약 (FA계약 시트)
    contracts = []
    for r in rows(wb['FA계약']):
        years_ = contract_years(r['기간(년)'])
        if not isinstance(r['계약 첫 시즌'], int) or not years_:
            continue
        contracts.append({
            'first': r['계약 첫 시즌'], 'id': str(r['선수ID']), 'name': r['선수명'],
            'kind': 'FA' if r['종류'] == 'FA' else 'nonFA', 'years': years_,
            'guaranteed': r['보장액(억 원)'] or 0, 'option': r['옵션(억 원)'] or 0,
            **({'note': str(r['기간(년)'])} if not isinstance(r['기간(년)'], int) else {}),
        })

    # ---- 순위표
    teams = defaultdict(list)
    for r in rows(wb['시즌순위']):
        teams[r['시즌']].append({
            'name': r['팀'], 'franchise': r['프랜차이즈'], 'league': r['리그'], 'rank': r['최종 순위(KBO 표기)'],
            'g': r['경기'], 'w': r['승'], 'l': r['패'], 't': r['무'], 'pct': r['승률(KBO 표기)'],
            'home': r['홈(승-무-패)'], 'away': r['방문(승-무-패)'],
        })

    # ---- 시즌 기록
    bat = defaultdict(list)
    pit = defaultdict(list)
    for r in rows(wb['타자시즌']):
        row = [str(r['선수ID']) if c == '선수ID' else (POS.get(r[c]) if c == '포지션(보조)' else r[c]) for c in BAT_COLS]
        bat[r['시즌']].append(row + [None] * (len(BAT_FIELDS) - len(BAT_COLS)))
    wpbk = defaultdict(int)
    for r in rows(wb['투수시즌']):
        wpbk[r['시즌']] += (r['WP'] or 0) + (r['BK'] or 0)
        row = [str(r['선수ID']) if c == '선수ID' else (POS.get(r[c]) if c == '역할(보조)' else r[c]) for c in PIT_COLS]
        pit[r['시즌']].append(row + [None] * (len(PIT_FIELDS) - len(PIT_COLS)))

    # ---- 하드힛 가공값으로 빈 기록 채우기
    parks, park_report = park_factors(teams)
    hardhit_report = fill_running(bat) + fill_gmli(pit) + fill_pitch_roles(pit) + hand_report + park_report
    wiki_records, wiki_record_n = fill_records_from_wiki(bat, pit, teams)
    wiki_fixes = apply_record_fixes(bat, pit, teams)

    # ---- 추가 명단 (CSV): 그 해 기록이 없는 선수를 기록 0인 행으로 넣는다. 능력은 직전 시즌 기록에서 나온다
    extra_ids = set()
    if EXTRA_ROSTER.exists():
        with EXTRA_ROSTER.open(encoding='utf-8-sig', newline='') as f:
            for r in csv.DictReader(f):
                y, pid, team = int(r['시즌']), str(r['선수ID']).strip(), r['구단'].strip()
                m = next((p for p in players if p['id'] == pid), None)
                if not m:
                    raise SystemExit(f'추가_명단.csv: 없는 선수 ID {pid} ({r.get("선수명")})')
                if team not in {t['name'] for t in teams.get(y, [])}:
                    raise SystemExit(f'추가_명단.csv: {y}년에 없는 구단 {team} ({r.get("선수명")})')
                is_pit = r['구분'].strip() == '투수'
                have = any(row[0] == pid for row in (pit[y] if is_pit else bat[y]))
                if have:
                    continue
                if is_pit:
                    role = m.get('pos') if m.get('pos') in ('SP', 'RP', 'CL') else None
                    pit[y].append([pid, team, role] + [0] * (len(PIT_FIELDS) - 4) + [None])
                else:
                    pos = m.get('pos') if m.get('pos') not in ('SP', 'RP', 'CL', 'P') else None
                    bat[y].append([pid, team, pos] + [0] * (len(BAT_COLS) - 3) + [None] * (len(BAT_FIELDS) - len(BAT_COLS)))
                extra_ids.add((y, pid))
                if m.get('last', 0) < y:
                    m['last'] = y

    # ---- 하드힛 주 포지션 (추가 명단 행까지 채운다)
    pos_report, pos_unresolved = fill_positions(bat, kind_of)
    hardhit_report += pos_report
    write_no_position(pos_unresolved, bat, players)

    years = sorted(teams)
    bi = {f: i for i, f in enumerate(BAT_FIELDS)}
    pi = {f: i for i, f in enumerate(PIT_FIELDS)}

    def n(v):
        return v or 0

    # ---- 리그 합계 (타자 쪽: 타석·안타 종류, 투수 쪽: 삼진·볼넷·사구·피홈런. 과거 타자 기록 일부에 볼넷·삼진이 비어 있어 투수 쪽을 쓴다)
    league = {}
    for y in years:
        b, p = bat[y], pit[y]
        L = {
            'teamGames': sum(t['g'] for t in teams[y]),
            'pa': sum(n(r[bi['pa']]) for r in b), 'ab': sum(n(r[bi['ab']]) for r in b), 'h': sum(n(r[bi['h']]) for r in b),
            'd': sum(n(r[bi['d']]) for r in b), 't': sum(n(r[bi['t']]) for r in b), 'hr': sum(n(r[bi['hr']]) for r in b),
            'r': sum(n(r[bi['r']]) for r in b), 'gdp': sum(n(r[bi['gdp']]) for r in b),
            'sb': sum(n(r[bi['sb']]) for r in b), 'sf': sum(n(r[bi['sf']]) for r in b), 'sac': sum(n(r[bi['sac']]) for r in b),
            'tbf': sum(n(r[pi['tbf']]) for r in p), 'bb': sum(n(r[pi['bb']]) for r in p), 'hbp': sum(n(r[pi['hbp']]) for r in p),
            'so': sum(n(r[pi['so']]) for r in p), 'hrAllowed': sum(n(r[pi['hr']]) for r in p),
            'wpbk': wpbk.get(y, 0),
            'er': sum(n(r[pi['er']]) for r in p), 'runsAllowed': sum(n(r[pi['r']]) for r in p), 'outs': sum(n(r[pi['outs']]) for r in p),
        }
        pa, tbf = L['pa'], L['tbf']
        L['rates'] = {
            's1': (L['h'] - L['d'] - L['t'] - L['hr']) / pa, 'd2': L['d'] / pa, 't3': L['t'] / pa, 'hr': L['hr'] / pa,
            'bb': L['bb'] / tbf, 'hbp': L['hbp'] / tbf, 'so': L['so'] / tbf,
        }
        L['avg'] = L['h'] / L['ab']
        L['runsPerTeamGame'] = L['r'] / L['teamGames']
        L['era'] = L['er'] * 27 / L['outs']
        L['gdpPerPa'] = L['gdp'] / pa
        L['sacPerPa'] = L['sac'] / pa
        L['sfPerPa'] = L['sf'] / pa
        L['unearnedShare'] = 1 - L['er'] / L['runsAllowed']
        L['wpbkPerPa'] = L['wpbk'] / tbf if L['wpbk'] else None  # 폭투+보크. 2001년부터 기록
        league[y] = L
    park_ev, park_ev_report = park_events(parks, league)
    hardhit_report += park_ev_report

    # ---- 능력 산출용 사전 평균: 출전이 적은 선수-시즌의 리그 대비 비율 (전 시즌 합산)
    ev_b = {'s1': None, 'd2': 'd', 't3': 't', 'hr': 'hr', 'bb': 'bb', 'hbp': 'hbp', 'so': 'so'}
    num_b = defaultdict(float); den_b = defaultdict(float)
    num_p = defaultdict(float); den_p = defaultdict(float)
    for y in years:
        if y == years[-1]:
            continue  # 진행 중 시즌 제외
        R = league[y]['rates']
        for r in bat[y]:
            pa = n(r[bi['pa']])
            if kind_of.get(r[bi['id']]) == 'P' or not (0 < pa < 100):
                continue
            h, d, t, hr = n(r[bi['h']]), n(r[bi['d']]), n(r[bi['t']]), n(r[bi['hr']])
            cnt = {'s1': h - d - t - hr, 'd2': d, 't3': t, 'hr': hr}
            for e in ('s1', 'd2', 't3', 'hr'):
                num_b[e] += cnt[e]; den_b[e] += pa * R[e]
            if r[bi['bb']] is not None and r[bi['so']] is not None:
                for e in ('bb', 'hbp', 'so'):
                    num_b[e] += n(r[bi[ev_b[e]]]); den_b[e] += pa * R[e]
        for r in pit[y]:
            tbf = n(r[pi['tbf']])
            if not (0 < tbf < 150):
                continue
            hit_rate = R['s1'] + R['d2'] + R['t3']
            num_p['hit'] += n(r[pi['h']]) - n(r[pi['hr']]); den_p['hit'] += tbf * hit_rate
            for e in ('hr', 'bb', 'hbp', 'so'):
                num_p[e] += n(r[pi[e]]); den_p[e] += tbf * R[e]
    priors = {
        'batter': {e: round(num_b[e] / den_b[e], 4) for e in num_b},
        'pitcher': {e: round(num_p[e] / den_p[e], 4) for e in num_p},
        'note': '타자 100타석 미만, 투수 상대 타자 150명 미만 선수-시즌의 리그 대비 비율',
    }

    # ---- 특별 엔트리 (은퇴식 등으로 하루 1군에 등록된 선수)
    # 그 해 1~2경기, 3타석·상대 타자 3명 이하, 34세 이상, 그 해가 마지막 1군 시즌인 선수 (규칙으로 고른 값).
    # 새 게임은 시작 연도 명단의 선수를 리그에 넣지 않는다 (현역이 아니다). 진행 중 시즌은 "마지막 시즌" 조건이 늘 참이다.
    bi, pi = BAT_FIELDS.index, PIT_FIELDS.index
    pmap = {p['id']: p for p in players}
    special = {}
    for y in years:
        use = defaultdict(lambda: [0, 0, 0])  # 경기, 타석, 상대 타자
        for r in bat[y]:
            u = use[r[bi('id')]]
            u[0] = max(u[0], r[bi('g')] or 0)
            u[1] += r[bi('pa')] or 0
        for r in pit[y]:
            u = use[r[pi('id')]]
            u[0] = max(u[0], r[pi('g')] or 0)
            u[2] += r[pi('tbf')] or 0
        ids = []
        for pid, (g, pa, bf) in use.items():
            m = pmap.get(pid)
            if not m or not m.get('birthYear') or (y, pid) in extra_ids:
                continue  # 추가 명단의 기록 0 행은 은퇴식이 아니다
            if g <= 2 and pa <= 3 and bf <= 3 and y - m['birthYear'] >= 34 and m.get('last') == y:
                ids.append(pid)
        if ids:
            special[str(y)] = sorted(ids)

    # ---- 쓰기
    (OUT / 'seasons').mkdir(parents=True, exist_ok=True)
    dump = lambda o, p: p.write_text(json.dumps(o, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    dump({'version': 1, 'players': players}, OUT / 'players.json')
    dump_traits(build_traits({p['id'] for p in players}))
    drafts = dump_drafts(drafts, players)
    dump({'version': 1, 'note': '보조자료의 FA 계약 목록. 금액은 구단 발표와 대조하지 않았다. 옵션 제외 보장액(억 원)', 'contracts': contracts},
         OUT / 'contracts.json')
    for y in years:
        dump({
            'year': y, 'complete': y != years[-1], 'games': max(t['g'] for t in teams[y]) if y != years[-1] else 144,
            'teams': teams[y], 'league': league[y],
            'batFields': BAT_FIELDS, 'bat': bat[y], 'pitFields': PIT_FIELDS, 'pit': pit[y],
        }, OUT / 'seasons' / f'{y}.json')
    franchises = defaultdict(list)
    for y in years:
        for t in teams[y]:
            f = franchises[t['franchise']]
            if not f or f[-1]['name'] != t['name']:
                f.append({'name': t['name'], 'from': y, 'to': y})
            else:
                f[-1]['to'] = y
    dump({
        'version': 1, 'source': SRC.name, 'years': years, 'inProgress': [years[-1]],
        'franchises': franchises, 'successors': {'쌍방울': 'SSG', '현대계': '히어로즈'},
        'rules': {y: rules_for(y) for y in years}, 'priors': priors,
        **({'foreignOpening': foreign_opening} if foreign_opening else {}),
        'specialEntries': special,
        **({'parkFactors': parks} if parks else {}),
        **({'parkEvents': park_ev} if park_ev else {}),
    }, OUT / 'meta.json')
    wiki_mil = build_military(profiles_of(wb), OUT)
    build_managers(OUT)
    build_rule_changes(OUT)
    build_homes()
    build_schedule()
    build_ppa()
    build_running()
    build_plate()
    build_history(OUT)
    build_defense(OUT)
    build_pitches(OUT)
    if not NO_WIKI:
        write_wiki_report(wiki_players, wiki_records, wiki_record_n, wiki_mil, special, wiki_fixes)
    if hardhit_report:
        head = ['# 하드힛 가공값 반영 결과', '', '`python tools/build_data.py`가 만든다. 기록 일부: 하드힛(https://hardhit.ai), 비상업적 이용.', '']
        HARDHIT_REPORT.write_text('\n'.join(head + hardhit_report) + '\n', encoding='utf-8')
    total = sum(f.stat().st_size for f in OUT.rglob('*.json'))
    print(f'선수 {len(players):,}명, 시즌 {len(years)}개({years[0]}~{years[-1]}), '
          f'타자 시즌 {sum(len(v) for v in bat.values()):,}행, 투수 시즌 {sum(len(v) for v in pit.values()):,}행')
    print(f'계약 {len(contracts)}건, 학력 구분 {sum(1 for p in players if p.get("school"))}명')
    print(f'신인 지명: {len(drafts)}개 연도 {sum(len(v) for v in drafts.values()):,}명 ({min(drafts, default="-")}~{max(drafts, default="-")} 입단)')
    print(f'추가 명단: {len(extra_ids)}명')
    print(f'외국인 개막 명단: {", ".join(f"{y}년 {sum(len(v) for v in t.values())}명" for y, t in foreign_opening.items()) or "없음"}, 아시아쿼터 {len(asia_ids)}명')
    print(f'사전 평균(타자): {priors["batter"]}')
    print(f'사전 평균(투수): {priors["pitcher"]}')
    print(f'출력: {OUT} ({total / 1e6:.1f} MB)')


if __name__ == '__main__':
    main()
