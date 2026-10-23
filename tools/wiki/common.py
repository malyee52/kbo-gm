# -*- coding: utf-8 -*-
"""위키 자료 수집 하네스 공통 부분: 경로, 예의 바른 HTTP 요청, 캐시, 우리 자료 읽기.

Wikimedia 규칙(로봇 정책·User-Agent 정책·API 에티켓, 2026-10-23 확인)을 여기 상수로 박아 둔다.
명령줄이나 환경 변수로 바꿀 수 없게 일부러 설정으로 빼지 않았다.
  - User-Agent: 저장소 주소를 연락처로 (메일은 넣지 않음)
  - 직렬 요청만 (스레드·동시 요청 없음)
  - 요청 사이 1초 이상, 1초 넘게 걸린 응답 뒤에는 5초
  - Action API에는 maxlag=5, maxlag·ratelimited·429·503이면 지수 백오프
  - Action API의 titles=는 50개씩 묶는다
받은 원문은 저장소 밖(RAW)에 캐시하고, 다시 돌리면 받은 것은 건너뛴다.
"""
import csv
import gzip
import hashlib
import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent.parent
DATA = ROOT / 'public' / 'data'
OUT_DIR = ROOT / 'data-src' / 'wiki'
XLSX = ROOT / 'data-src' / 'KBO_단장게임_자료집.xlsx'

# 저장소 밖: 원문 캐시·덤프·정리 노트 (커밋하지 않는다)
RAW = Path.home() / 'kbo-gm-raw'
RAW_WIKI = RAW / 'wiki'
RAW_NAMU = RAW / 'namuwiki'
RAW_NOTE = RAW / '정리'
# 위키 값을 넣지 않은 기반 자료 (build_data.py --no-wiki). 5단계가 '우리 값'으로 견준다
BASE_DATA = RAW_NOTE / '기반'

# ---- Wikimedia 규칙 (고치지 말 것)
UA = f'kbo-gm-data/1.0 (https://github.com/malyee52/kbo-gm) python-urllib/{urllib.request.__version__}'
MIN_GAP = 1.0          # 요청 사이 최소 간격(초)
SLOW_RESPONSE = 1.0    # 이보다 오래 걸린 응답 뒤에는
SLOW_GAP = 5.0         # 이만큼 쉰다
MAXLAG = 5
BATCH = 50             # titles= 묶음 크기
MAX_TRIES = 7          # 백오프 재시도 (5, 10, 20, 40, 80, 160초)
KOWIKI_API = 'https://ko.wikipedia.org/w/api.php'
SPARQL = 'https://query.wikidata.org/sparql'

TODAY = date.today().isoformat()


def log(*a):
    print(*a, flush=True)


class Polite:
    """한 번에 하나씩, 간격을 지키며 요청한다. 프로세스 안에서 하나만 쓴다."""

    def __init__(self):
        self.next_ok = 0.0
        self.count = 0
        RAW_WIKI.mkdir(parents=True, exist_ok=True)
        self.logf = (RAW_WIKI / '요청기록.log').open('a', encoding='utf-8')

    def _wait(self):
        d = self.next_ok - time.monotonic()
        if d > 0:
            time.sleep(d)

    def get(self, url, params, post=False, action_api=False):
        """JSON 응답을 돌려준다. 실패가 계속되면 예외."""
        params = dict(params)
        if action_api:
            params.setdefault('format', 'json')
            params.setdefault('formatversion', '2')
            params['maxlag'] = MAXLAG
        body = urllib.parse.urlencode(params).encode('utf-8')
        delay = 5.0
        for attempt in range(1, MAX_TRIES + 1):
            self._wait()
            if post:
                req = urllib.request.Request(url, data=body, method='POST')
                req.add_header('Content-Type', 'application/x-www-form-urlencoded')
            else:
                req = urllib.request.Request(url + '?' + body.decode('ascii'))
            req.add_header('User-Agent', UA)
            req.add_header('Accept-Encoding', 'gzip')
            if url == SPARQL:
                req.add_header('Accept', 'application/sparql-results+json')
            t0 = time.monotonic()
            retry_after = None
            err = None
            data = None
            try:
                with urllib.request.urlopen(req, timeout=120) as r:
                    raw = r.read()
                    if r.headers.get('Content-Encoding') == 'gzip':
                        raw = gzip.decompress(raw)
                    data = json.loads(raw.decode('utf-8'))
            except urllib.error.HTTPError as e:
                err = f'HTTP {e.code}'
                if e.code in (429, 503, 502, 504):
                    retry_after = e.headers.get('Retry-After')
                else:
                    self._done(t0, url, err)
                    raise
            except (urllib.error.URLError, TimeoutError, ConnectionError) as e:
                err = f'연결 오류 {e}'
            self._done(t0, url, err or 'ok')
            if data is not None and action_api and 'error' in data:
                code = data['error'].get('code')
                if code in ('maxlag', 'ratelimited'):
                    err = f'API {code}'
                    data = None
                else:
                    raise RuntimeError(f'API 오류: {data["error"]}')
            if data is not None:
                return data
            wait = delay
            if retry_after and retry_after.isdigit():
                wait = max(wait, float(retry_after))
            log(f'  ! {err} → {wait:.0f}초 뒤 다시 ({attempt}/{MAX_TRIES})')
            time.sleep(wait)
            delay *= 2
        raise RuntimeError(f'요청이 계속 실패: {url} ({err})')

    def _done(self, t0, url, status):
        took = time.monotonic() - t0
        self.count += 1
        gap = SLOW_GAP if took > SLOW_RESPONSE else MIN_GAP
        self.next_ok = time.monotonic() + gap
        self.logf.write(f'{time.strftime("%Y-%m-%d %H:%M:%S")}\t{url}\t{status}\t{took:.2f}s\n')
        self.logf.flush()


# ---- 문서 캐시 (한국어 위키백과)

PAGES = RAW_WIKI / 'pages'


def _page_file(title):
    return PAGES / (hashlib.sha1(title.encode('utf-8')).hexdigest() + '.json')


def cached_page(title):
    f = _page_file(title)
    if f.exists():
        return json.loads(f.read_text(encoding='utf-8'))
    return None


def fetch_pages(http, titles, label='문서'):
    """제목들의 최신 원문을 캐시에서 주거나 50개씩 받아 캐시한다.
    돌려주는 값: 요청 제목 → 기록 {requested, title, missing, pageid, revid, timestamp, qid, disambig, content, fetched}"""
    PAGES.mkdir(parents=True, exist_ok=True)
    out = {}
    todo = []
    for t in dict.fromkeys(titles):
        c = cached_page(t)
        if c is not None:
            out[t] = c
        else:
            todo.append(t)
    if todo:
        log(f'  {label}: 캐시 {len(out)}쪽, 새로 받을 제목 {len(todo)}개 ({(len(todo) + BATCH - 1) // BATCH}번 요청 + 이어 받기)')
    for i in range(0, len(todo), BATCH):
        chunk = todo[i:i + BATCH]
        got = _query_pages(http, chunk)
        for t in chunk:
            rec = got[t]
            _page_file(t).write_text(json.dumps(rec, ensure_ascii=False), encoding='utf-8')
            out[t] = rec
        log(f'  {label}: {min(i + BATCH, len(todo))}/{len(todo)}')
    return out


def _query_pages(http, chunk):
    params = {
        'action': 'query', 'prop': 'revisions|pageprops', 'rvprop': 'content|ids|timestamp', 'rvslots': 'main',
        'ppprop': 'wikibase_item|disambiguation', 'redirects': '1', 'titles': '|'.join(chunk),
    }
    pages = {}
    norm = {}
    cont = {}
    while True:
        d = http.get(KOWIKI_API, {**params, **cont}, post=True, action_api=True)
        q = d.get('query', {})
        for n in q.get('normalized', []):
            norm[n['from']] = n['to']
        for n in q.get('redirects', []):
            norm[n['from']] = n['to']
        for p in q.get('pages', []):
            prev = pages.get(p['title'], {})
            revs = p.get('revisions') or prev.get('revisions')
            pages[p['title']] = {**prev, **p, 'revisions': revs}
        if 'continue' not in d:
            break
        cont = d['continue']
    res = {}
    for t in chunk:
        final = t
        seen = set()
        while final in norm and final not in seen:
            seen.add(final)
            final = norm[final]
        p = pages.get(final)
        rec = {'requested': t, 'title': final, 'fetched': TODAY}
        if p is None or p.get('missing') or p.get('invalid'):
            rec['missing'] = True
        else:
            rev = (p.get('revisions') or [{}])[0]
            pp = p.get('pageprops') or {}
            rec.update({
                'pageid': p.get('pageid'), 'revid': rev.get('revid'), 'timestamp': rev.get('timestamp'),
                'qid': pp.get('wikibase_item'), 'disambig': 'disambiguation' in pp,
                'content': (rev.get('slots') or {}).get('main', {}).get('content'),
            })
        res[t] = rec
    return res


# ---- 우리 자료

def load_players(data=None):
    return json.loads(((data or DATA) / 'players.json').read_text(encoding='utf-8'))['players']


def load_meta_years(data=None):
    return sorted(int(p.stem) for p in ((data or DATA) / 'seasons').glob('*.json'))


def load_seasons(data=None):
    """연도 → 시즌 JSON"""
    data = data or DATA
    return {y: json.loads((data / 'seasons' / f'{y}.json').read_text(encoding='utf-8')) for y in load_meta_years(data)}


def build_base_data():
    """위키 값을 넣지 않은 기반 자료를 BASE_DATA에 만든다 (build_data.py --no-wiki, 1분 안팎). 경로를 돌려준다"""
    import shutil
    import subprocess
    BASE_DATA.mkdir(parents=True, exist_ok=True)
    if not (BASE_DATA / 'meta.json').exists():
        # build_data.py의 지명 단계가 출력 폴더의 앞 빌드(meta·seasons)에서 통산 경기를 읽으므로 처음에는 게임 자료를 옮겨 둔다
        shutil.copy(DATA / 'meta.json', BASE_DATA / 'meta.json')
        shutil.copytree(DATA / 'seasons', BASE_DATA / 'seasons', dirs_exist_ok=True)
    subprocess.run([sys.executable, str(ROOT / 'tools' / 'build_data.py'), str(XLSX), str(BASE_DATA), '--no-wiki'],
                   check=True, stdout=subprocess.DEVNULL)
    return BASE_DATA


def player_stints(seasons):
    """선수 ID → {(연도, 계보)}  (1군 기록 행 기준)"""
    fr = {}
    for y, s in seasons.items():
        for t in s['teams']:
            fr[(y, t['name'])] = t['franchise']
    out = {}
    for y, s in seasons.items():
        for key in ('bat', 'pit'):
            for r in s[key]:
                out.setdefault(r[0], set()).add((y, fr.get((y, r[1]), r[1])))
    return out


def load_master_extra():
    """자료집 선수마스터에서 players.json에 없는 칸(생년월일, 신장/체중, 출신교·경력)을 읽는다."""
    from openpyxl import load_workbook
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb['선수마스터']
    it = ws.iter_rows(values_only=True)
    head = next(it)
    ix = {h: i for i, h in enumerate(head)}
    out = {}
    for r in it:
        if r[ix['선수ID']] is None:
            continue
        out[str(r[ix['선수ID']])] = {
            'birth': r[ix['생년월일']], 'hw': r[ix['신장/체중(2026 프로필)']],
            'career': r[ix['출신교·경력(2026 프로필)']], 'hand': r[ix['투타']], 'pos': r[ix['주포지션']],
        }
    wb.close()
    return out


# 구단 이름 조각 → 계보 (시즌 JSON의 franchise 값)
FRANCHISE_KEYS = [
    ('MBC', 'LG'), ('LG', 'LG'), ('삼성', '삼성'), ('롯데', '롯데'), ('해태', 'KIA'), ('KIA', 'KIA'), ('기아', 'KIA'),
    ('OB', '두산'), ('두산', '두산'), ('삼미', '현대계'), ('청보', '현대계'), ('태평양', '현대계'), ('현대', '현대계'),
    ('빙그레', '한화'), ('한화', '한화'), ('쌍방울', '쌍방울'), ('SSG', 'SSG'), ('SK', 'SSG'), ('우리', '히어로즈'),
    ('서울 히어로즈', '히어로즈'), ('히어로즈', '히어로즈'), ('넥센', '히어로즈'), ('키움', '히어로즈'), ('NC', 'NC'),
    ('KT', 'KT'), ('kt', 'KT'),
]


def franchise_of(text):
    """구단 이름(글)에서 계보를 찾는다. 못 찾으면 None"""
    if not text:
        return None
    for k, v in FRANCHISE_KEYS:
        if re.search(r'(?<![A-Za-z])' + re.escape(k) + r'(?![A-Za-z])', text):
            return v
    return None


# 시즌 기록 표의 구단 칸 읽기 (2026-10-23 고침). franchise_of는 글자 조각으로 찾아서
# '요미우리'의 '우리'(→ 히어로즈), '지바 롯데'(→ 롯데)처럼 해외 구단을 KBO 구단으로 읽었다.
# 여기서는 칸을 낱말로 나눠 낱말이 모두 KBO 구단 이름·별칭이어야 KBO 구단으로 본다.
KBO_TEAM_WORDS = {
    'MBC': 'LG', 'LG': 'LG', '삼성': '삼성', '롯데': '롯데', '해태': 'KIA', 'KIA': 'KIA', '기아': 'KIA',
    'OB': '두산', '두산': '두산', '삼미': '현대계', '청보': '현대계', '태평양': '현대계', '현대': '현대계',
    '빙그레': '한화', '한화': '한화', '쌍방울': '쌍방울', 'SK': 'SSG', 'SSG': 'SSG', '우리': '히어로즈',
    '히어로즈': '히어로즈', '넥센': '히어로즈', '키움': '히어로즈', 'NC': 'NC', 'KT': 'KT',
}
# 별칭 꼬리·지역명: 낱말로 나와도 되지만 계보를 정하지 않는다
TEAM_NICKS = {
    '청룡', '트윈스', '라이온즈', '자이언츠', '타이거즈', '베어스', '슈퍼스타즈', '핀토스', '돌핀스', '유니콘스', '이글스',
    '레이더스', '와이번스', '랜더스', '다이노스', '위즈', 'WIZ', '서울', '서을',  # '서을': 나무위키 '우리 서을 넥센' 오타
}
# 군 야구단(퓨처스): 칸에 같이 적혀 있으면 그 조각만 버린다 ('kt(~6.3) 상무(7.19~)')
FARM_WORDS = {'상무', '경찰', '경찰청', '국군체육부대', '상무피닉스'}


def franchises_of(text):
    """시즌 기록 표의 구단 칸 → KBO 계보 목록 (시즌 중 이적 'LG/롯데'는 둘, 칸에 적힌 순서).
    KBO 구단이 아닌 낱말(요미우리·지바·컵스·2군 …)이 하나라도 있으면 None, 군 야구단뿐이면 []."""
    t = re.sub(r'\[[^\]]*\]', ' ', text or '')     # 각주 [^8]
    t = re.sub(r'\([^)]*\)', ' ', t)               # (~6.3) 같은 날짜
    out = []
    for w in re.split(r'[\s/,·→\-–~～]+', t):
        if not w:
            continue
        u = w.upper()
        if u in KBO_TEAM_WORDS:
            fr = KBO_TEAM_WORDS[u]
            if fr not in out:
                out.append(fr)
        elif w in KBO_TEAM_WORDS:
            if KBO_TEAM_WORDS[w] not in out:
                out.append(KBO_TEAM_WORDS[w])
        elif w in TEAM_NICKS or u in TEAM_NICKS or w in FARM_WORDS:
            continue
        else:
            return None
    return out


def write_csv(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(header)
        w.writerows(rows)


def read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def strip_paren(title):
    """'김경남 (야구 선수)' → '김경남'"""
    return re.sub(r'\s*\([^)]*\)\s*$', '', title or '').strip()


def sample_ids():
    f = Path(__file__).with_name('sample_ids.txt')
    return [l.split('#')[0].strip() for l in f.read_text(encoding='utf-8').splitlines() if l.split('#')[0].strip()]


def out_paths(sample):
    """견본 실행은 저장소 밖 견본 폴더에, 전체 실행은 data-src/wiki/에 쓴다."""
    if sample:
        return RAW_NOTE / '견본' / 'data-src-wiki', RAW_NOTE / '견본'
    return OUT_DIR, RAW_NOTE


def stdout_utf8():
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass
